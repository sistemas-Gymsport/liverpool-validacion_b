"""Verificador de pedidos Liverpool contra picking de Odoo y guías de envío.

Primera revisión (Excel): reporte de pedidos de Liverpool contra la tabla de medidas de equipos.
Segunda revisión (PDF): picking de Odoo y guías de envío contra lo anterior.

Uso por línea de comandos (los documentos se identifican solos):
    python verificar_pedidos.py "reporte de pedidos.xlsx" [medidas.xlsx] [picking.pdf ...] [guias.pdf ...] [-o reporte.xlsx]

Funciones públicas:
    analizar(documentos) -> Resultado        documentos: lista de (nombre, archivo) en cualquier orden
    verificar(excel, pdfs, medidas=None, guias=None) -> Resultado
    escribir_reporte(resultado, destino) -> None
"""

from __future__ import annotations

import argparse
import io
import logging
import re
import sys
import unicodedata
from dataclasses import dataclass, field
from datetime import datetime
from difflib import SequenceMatcher
from pathlib import Path
from typing import BinaryIO, Iterable, Union

import pdfplumber
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pypdf import PdfReader

import guias_envio as ge
import validar_guias as vg

Fuente = Union[str, Path, bytes, BinaryIO]

ERROR = "ERROR"
REVISAR = "REVISAR"
CORRECTO = "CORRECTO"

ORIGEN_EXCEL = "Excel"
ORIGEN_PICKING = "Picking"
ORIGEN_GUIAS = "Guías de envío"

TIPO_LIVERPOOL = "Reporte de pedidos Liverpool"
TIPO_MEDIDAS = "Tabla de medidas de equipos"
TIPO_PICKING = "Picking de Odoo"
TIPO_GUIAS = "Guías de envío"
TIPO_DESCONOCIDO = "No reconocido"

ESTADO_ESPERADO = "Pendiente de envío"
MAX_DIGITOS_DIFERENTES = 2
SIMILITUD_NOMBRE_SUGERENCIA = 0.6
TOLERANCIA_RENGLON = 6.0

COLUMNAS_OBLIGATORIAS = ("Id del pedido", "SKU de la oferta", "Cantidad")
COLUMNAS_OPCIONALES = (
    "Estado",
    "Nombre completo del cliente que recibe",
    "Nombre del producto",
    "Fecha límite de envío",
    "Método de envío",
    "No de guía de entrega",
    "Estado de solicitud de la guía",
    "Reembolso total",
    "Reembolso parcial",
    "Pedidos con reembolso",
    "Largo del paquete (s)",
    "Alto del paquete (s)",
    "Ancho del paquete (s)",
    "Peso del paquete (s)",
    "No de guías requeridas para el envío",
)

RE_PEDIDO = re.compile(r"^\s*(\d{6,14})\s*/\s*(.*)$")
RE_PIEZAS = re.compile(r"^\s*(\d+)\s*(piezas?|pzas?\.?|pz\.?|pares?|unidades?|uds?\.?)\s*$", re.I)
RE_CANTIDAD = re.compile(r"^\d+(?:[.,]\d+)?$")
RE_PICKING = re.compile(r"\b[A-Z0-9]+/PICK/\d+\b")
ABREVIATURAS = {"ma": "maria", "ma.": "maria", "ma.del": "maria del", "fco": "francisco", "jse": "jose"}

logging.getLogger("pypdf").setLevel(logging.ERROR)
logging.getLogger("pdfminer").setLevel(logging.ERROR)


class DocumentoInvalido(ValueError):
    """El documento no tiene el formato esperado."""


@dataclass
class LineaLiverpool:
    pedido: str
    sku: str
    cantidad: int
    estado: str
    cliente: str
    producto: str
    fecha_limite: str
    envio: str
    guia: str
    estado_guia: str
    reembolsado: bool
    paquetes: vg.Paquetes = field(default_factory=lambda: vg.Paquetes(None, [], ""))


@dataclass
class Anotacion:
    texto: str
    pedido: str
    nombre: str
    piezas: int | None
    pagina: int
    centro_y: float


@dataclass
class RenglonPicking:
    picking: str
    archivo: str
    pagina: int
    sku: str
    cantidad: float
    unidad: str
    producto: str
    top: float
    anotaciones: list[Anotacion] = field(default_factory=list)


@dataclass
class Picking:
    nombre: str
    archivo: str
    renglones: list[RenglonPicking]
    anotaciones_huerfanas: list[Anotacion]
    total_anotaciones: int


@dataclass
class Hallazgo:
    nivel: str
    picking: str
    pedido: str
    sku: str
    problema: str
    origen: str = ORIGEN_PICKING


@dataclass
class RevisionEtiqueta:
    resultado: str
    archivo: str
    pagina: int
    paqueteria: str
    rastreo: str
    pedido: str
    identificacion: str
    contenido: str
    peso: str
    peso_estimado: str
    medidas: str
    observaciones: list[str] = field(default_factory=list)


@dataclass
class DocumentoAnalizado:
    nombre: str
    tipo: str
    detalle: str


@dataclass
class DetallePedido:
    resultado: str
    picking: str
    pagina: int
    pedido: str
    sku_pdf: str
    cantidad_renglon: float
    anotacion: str
    cliente_liverpool: str
    sku_liverpool: str
    cantidad_liverpool: str
    envio: str
    guia: str
    productos_pedido: int = 1
    observaciones: list[str] = field(default_factory=list)


@dataclass
class SinPicking:
    pedido: str
    sku: str
    cantidad: int
    cliente: str
    producto: str
    estado: str
    fecha_limite: str
    envio: str
    observacion: str


@dataclass
class Resultado:
    pedidos_liverpool: int
    verificados: int
    correctos: int
    errores: int
    revisar: int
    sin_picking: int | None
    hallazgos: list[Hallazgo]
    detalle: list[DetallePedido]
    pedidos_sin_picking: list[SinPicking]
    advertencias: list[str]
    pickings: list[str]
    guias: list[vg.RevisionGuias] = field(default_factory=list)
    guias_por_revisar: int = 0
    medidas_cargadas: bool = False
    picking_cargado: bool = True
    etiquetas: list[RevisionEtiqueta] = field(default_factory=list)
    guias_envio_cargadas: bool = False
    documentos: list[DocumentoAnalizado] = field(default_factory=list)

    @property
    def todo_cuadra(self) -> bool:
        return self.errores == 0 and self.revisar == 0 and not self.advertencias


# ---------------------------------------------------------------------------
# Normalización
# ---------------------------------------------------------------------------

def normalizar_texto(valor: object) -> str:
    texto = "" if valor is None else str(valor)
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", texto).strip().lower()


def normalizar_sku(valor: object) -> str:
    return re.sub(r"\s+", "", "" if valor is None else str(valor)).upper()


def tokens_nombre(nombre: str) -> list[str]:
    texto = normalizar_texto(nombre)
    texto = re.sub(r"[^a-z0-9. ]", " ", texto)
    tokens = []
    for t in texto.split():
        tokens.extend(ABREVIATURAS.get(t, t.rstrip(".")).split())
    return [t for t in tokens if t]


def nombres_coinciden(anotado: str, liverpool: str) -> bool:
    a, b = tokens_nombre(anotado), tokens_nombre(liverpool)
    if not a or not b:
        return False
    if a == b:
        return True
    menor, mayor = (a, b) if len(a) <= len(b) else (b, a)
    return len(menor) >= 2 and set(menor) <= set(mayor)


def similitud_nombres(a: str, b: str) -> float:
    ta, tb = " ".join(sorted(tokens_nombre(a))), " ".join(sorted(tokens_nombre(b)))
    if not ta or not tb:
        return 0.0
    return SequenceMatcher(None, ta, tb).ratio()


def digitos_diferentes(a: str, b: str) -> int:
    if len(a) == len(b):
        return sum(1 for x, y in zip(a, b) if x != y)
    previo = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        actual = [i]
        for j, cb in enumerate(b, 1):
            actual.append(min(previo[j] + 1, actual[j - 1] + 1, previo[j - 1] + (ca != cb)))
        previo = actual
    return previo[-1]


def a_entero(valor: object) -> int:
    texto = str(valor or "").strip().replace(",", ".")
    try:
        return int(round(float(texto)))
    except ValueError:
        return 0


def formatear_cantidad(valor: float) -> str:
    return str(int(valor)) if float(valor).is_integer() else str(valor)


def es_si(valor: object) -> bool:
    return normalizar_texto(valor) in {"si", "sí", "yes", "true", "1"}


def parsear_fecha_limite(texto: str) -> datetime:
    limpio = normalizar_texto(texto).replace(".", "")
    for formato in ("%d/%m/%Y - %I:%M %p", "%d/%m/%Y %I:%M %p", "%d/%m/%Y - %H:%M", "%d/%m/%Y %H:%M", "%d/%m/%Y"):
        try:
            return datetime.strptime(limpio.upper(), formato)
        except ValueError:
            continue
    return datetime.max


def _abrir_binario(fuente: Fuente) -> BinaryIO:
    if isinstance(fuente, (str, Path)):
        return open(fuente, "rb")
    if isinstance(fuente, (bytes, bytearray)):
        return io.BytesIO(fuente)
    fuente.seek(0)
    return fuente


def _con_nombre(elemento, indice: int) -> tuple[str, Fuente]:
    if isinstance(elemento, tuple):
        return elemento
    return _nombre_fuente(elemento, indice), elemento


def _nombre_fuente(fuente: Fuente, indice: int) -> str:
    if isinstance(fuente, (str, Path)):
        return Path(fuente).name
    return getattr(fuente, "name", None) or f"PDF {indice + 1}"


# ---------------------------------------------------------------------------
# Excel Liverpool
# ---------------------------------------------------------------------------

def leer_liverpool(excel: Fuente) -> list[LineaLiverpool]:
    manejador = _abrir_binario(excel)
    try:
        try:
            libro = load_workbook(manejador, read_only=True, data_only=True)
        except Exception as exc:
            raise DocumentoInvalido("El archivo Excel no tiene un formato válido.") from exc
        try:
            hoja = libro["OrderReport"] if "OrderReport" in libro.sheetnames else libro.worksheets[0]
            hoja.reset_dimensions()
            filas = hoja.iter_rows(values_only=True)
            indices = _ubicar_encabezados(filas)
            lineas = [_linea_desde_fila(fila, indices) for fila in filas if _fila_con_pedido(fila, indices)]
        finally:
            libro.close()
    finally:
        if isinstance(excel, (str, Path)):
            manejador.close()
    return lineas


def _ubicar_encabezados(filas: Iterable[tuple]) -> dict[str, int]:
    buscadas = {normalizar_texto(c): c for c in COLUMNAS_OBLIGATORIAS + COLUMNAS_OPCIONALES}
    for numero, fila in enumerate(filas):
        if numero >= 15:
            break
        indices: dict[str, int] = {}
        for i, celda in enumerate(fila):
            original = buscadas.get(normalizar_texto(celda))
            if original and original not in indices:
                indices[original] = i
        if all(c in indices for c in COLUMNAS_OBLIGATORIAS):
            return indices
    raise DocumentoInvalido("El Excel no parece ser el reporte de pedidos de Liverpool.")


def _celda(fila: tuple, indices: dict[str, int], columna: str) -> str:
    i = indices.get(columna)
    if i is None or i >= len(fila) or fila[i] is None:
        return ""
    valor = fila[i]
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


def _fila_con_pedido(fila: tuple, indices: dict[str, int]) -> bool:
    return bool(_celda(fila, indices, "Id del pedido"))


def _linea_desde_fila(fila: tuple, indices: dict[str, int]) -> LineaLiverpool:
    c = lambda col: _celda(fila, indices, col)  # noqa: E731
    estado = c("Estado")
    reembolsado = (
        es_si(c("Reembolso total"))
        or es_si(c("Reembolso parcial"))
        or bool(c("Pedidos con reembolso"))
        or "reembols" in normalizar_texto(estado)
    )
    return LineaLiverpool(
        pedido=re.sub(r"\s+", "", c("Id del pedido")),
        sku=normalizar_sku(c("SKU de la oferta")),
        cantidad=a_entero(c("Cantidad")),
        estado=estado,
        cliente=c("Nombre completo del cliente que recibe"),
        producto=c("Nombre del producto"),
        fecha_limite=c("Fecha límite de envío"),
        envio=c("Método de envío"),
        guia=c("No de guía de entrega"),
        estado_guia=c("Estado de solicitud de la guía"),
        reembolsado=reembolsado,
        paquetes=vg.leer_paquetes(c("Largo del paquete (s)"), c("Alto del paquete (s)"), c("Ancho del paquete (s)"),
                                  c("Peso del paquete (s)"), c("No de guías requeridas para el envío")),
    )


# ---------------------------------------------------------------------------
# PDF Odoo
# ---------------------------------------------------------------------------

def leer_picking(pdf: Fuente, nombre_archivo: str) -> Picking:
    manejador = _abrir_binario(pdf)
    try:
        datos = manejador.read()
    finally:
        if isinstance(pdf, (str, Path)):
            manejador.close()
    try:
        lector = PdfReader(io.BytesIO(datos))
        if lector.is_encrypted:
            raise DocumentoInvalido(f"El documento PDF «{nombre_archivo}» está protegido con contraseña.")
        anotaciones_freetext = _anotaciones_freetext(lector)
        with pdfplumber.open(io.BytesIO(datos)) as documento:
            return _procesar_paginas(documento, nombre_archivo, anotaciones_freetext)
    except DocumentoInvalido:
        raise
    except Exception as exc:
        raise DocumentoInvalido(f"El documento PDF «{nombre_archivo}» está dañado o no puede leerse.") from exc


def _anotaciones_freetext(lector: PdfReader) -> dict[int, list[Anotacion]]:
    resultado: dict[int, list[Anotacion]] = {}
    for numero, pagina in enumerate(lector.pages):
        alto = float(pagina.mediabox.height)
        for ref in pagina.get("/Annots") or []:
            anotacion = ref.get_object()
            if anotacion.get("/Subtype") != "/FreeText":
                continue
            contenido = anotacion.get("/Contents")
            rect = anotacion.get("/Rect")
            if not contenido or not rect:
                continue
            y1, y2 = float(rect[1]), float(rect[3])
            centro = alto - (y1 + y2) / 2
            resultado.setdefault(numero, []).extend(_parsear_anotacion(str(contenido), numero + 1, centro))
    return resultado


def _parsear_anotacion(texto: str, pagina: int, centro_y: float) -> list[Anotacion]:
    anotaciones: list[Anotacion] = []
    for linea in re.split(r"[\r\n]+", texto):
        linea = linea.strip()
        if not linea:
            continue
        coincidencia = RE_PEDIDO.match(linea)
        if coincidencia:
            anotaciones.append(_crear_anotacion(linea, coincidencia, pagina, centro_y))
        elif anotaciones:
            previa = anotaciones[-1]
            texto_unido = f"{previa.texto} {linea}"
            anotaciones[-1] = _crear_anotacion(texto_unido, RE_PEDIDO.match(texto_unido), pagina, previa.centro_y)
    return anotaciones


def _crear_anotacion(texto: str, coincidencia: re.Match, pagina: int, centro_y: float) -> Anotacion:
    partes = [p.strip() for p in coincidencia.group(2).split("/")]
    nombre_partes, piezas = [], None
    for parte in partes:
        cantidad = RE_PIEZAS.match(parte)
        if cantidad:
            piezas = int(cantidad.group(1))
        elif parte:
            nombre_partes.append(parte)
    return Anotacion(
        texto=re.sub(r"\s+", " ", texto).strip(),
        pedido=coincidencia.group(1),
        nombre=" ".join(nombre_partes),
        piezas=piezas,
        pagina=pagina,
        centro_y=centro_y,
    )


def _familia_fuente(nombre: str) -> str:
    base = nombre.split("+", 1)[-1]
    return re.split(r"[-,]", base, 1)[0].lower()


def _procesar_paginas(documento, nombre_archivo: str, freetext: dict[int, list[Anotacion]]) -> Picking:
    nombre_picking = ""
    renglones: list[RenglonPicking] = []
    huerfanas: list[Anotacion] = []
    total_anotaciones = 0
    familia_odoo = None

    for numero, pagina in enumerate(documento.pages):
        palabras = pagina.extract_words(extra_attrs=["fontname", "size"], keep_blank_chars=False)
        if familia_odoo is None:
            familia_odoo = _detectar_familia_odoo(palabras)
        palabras_odoo = [p for p in palabras if _es_fuente_odoo(p["fontname"], familia_odoo)]
        palabras_anotacion = [p for p in palabras if not _es_fuente_odoo(p["fontname"], familia_odoo)]

        if not nombre_picking:
            texto = " ".join(p["text"] for p in palabras_odoo)
            encontrado = RE_PICKING.search(texto)
            nombre_picking = encontrado.group(0) if encontrado else ""

        renglones_pagina = _extraer_renglones(palabras_odoo, numero + 1, float(pagina.height))
        anotaciones = _anotaciones_aplanadas(palabras_anotacion, numero + 1)
        anotaciones += _sin_duplicar(freetext.get(numero, []), anotaciones)
        total_anotaciones += len(anotaciones)

        for anotacion in sorted(anotaciones, key=lambda a: a.centro_y):
            renglon = _renglon_para(anotacion.centro_y, renglones_pagina)
            if renglon is None:
                huerfanas.append(anotacion)
            else:
                renglon.anotaciones.append(anotacion)
        renglones.extend(r for r, _ in renglones_pagina)

    nombre = nombre_picking or Path(nombre_archivo).stem
    for renglon in renglones:
        renglon.picking = nombre
    return Picking(nombre, nombre_archivo, renglones, huerfanas, total_anotaciones)


def _detectar_familia_odoo(palabras: list[dict]) -> str | None:
    for palabra in palabras:
        if palabra["text"].startswith("[") or palabra["text"].upper() in {"PRODUCTO", "CANTIDAD"}:
            return _familia_fuente(palabra["fontname"])
    return None


def _es_fuente_odoo(fuente: str, familia: str | None) -> bool:
    if familia is None:
        return True
    familia_palabra = _familia_fuente(fuente)
    return familia_palabra == familia or familia_palabra == "fontawesome"


def _columnas(palabras: list[dict]) -> tuple[float, float] | None:
    cantidad = next((p for p in palabras if p["text"].upper() == "CANTIDAD"), None)
    desde = next((p for p in palabras if p["text"].upper() == "DESDE" and cantidad and abs(p["top"] - cantidad["top"]) < 3), None)
    if not cantidad:
        return None
    return cantidad["x0"] - 12, (desde["x0"] - 2 if desde else cantidad["x1"] + 40)


def _extraer_renglones(palabras: list[dict], pagina: int, alto: float) -> list[tuple[RenglonPicking, float]]:
    columnas = _columnas(palabras)
    if columnas is None:
        return []
    x_cant_ini, x_cant_fin = columnas
    encabezado = next(p for p in palabras if p["text"].upper() == "CANTIDAD")
    lineas = _agrupar_lineas([p for p in palabras if p["top"] > encabezado["bottom"]])
    pie = _inicio_pie(lineas, alto)

    inicios = []
    for i, linea in enumerate(lineas):
        primera = linea[0]
        if primera["text"].startswith("[") and primera["x0"] < x_cant_ini:
            inicios.append(i)

    resultado = []
    for n, inicio in enumerate(inicios):
        fin = inicios[n + 1] if n + 1 < len(inicios) else len(lineas)
        bloque = [l for l in lineas[inicio:fin] if l[0]["top"] < pie]
        sku, resto_producto = _sku_de_bloque(bloque, x_cant_ini)
        cantidad, unidad = _cantidad_de_bloque(bloque, x_cant_ini, x_cant_fin)
        producto = " ".join(p["text"] for l in bloque for p in l if p["x0"] < x_cant_ini and p not in resto_producto)
        top = lineas[inicio][0]["top"]
        siguiente = lineas[inicios[n + 1]][0]["top"] if n + 1 < len(inicios) else pie
        resultado.append((RenglonPicking("", "", pagina, sku, cantidad, unidad, producto.strip(), top), siguiente))
    return resultado


def _agrupar_lineas(palabras: list[dict]) -> list[list[dict]]:
    lineas: list[list[dict]] = []
    for palabra in sorted(palabras, key=lambda p: (round(p["top"]), p["x0"])):
        if lineas and abs(lineas[-1][0]["top"] - palabra["top"]) < 3:
            lineas[-1].append(palabra)
        else:
            lineas.append([palabra])
    for linea in lineas:
        linea.sort(key=lambda p: p["x0"])
    return lineas


def _inicio_pie(lineas: list[list[dict]], alto: float) -> float:
    for linea in lineas:
        if linea[0]["top"] > alto * 0.85:
            return linea[0]["top"] - 1
    return alto


def _sku_de_bloque(bloque: list[list[dict]], x_limite: float) -> tuple[str, list[dict]]:
    usadas: list[dict] = []
    texto = ""
    for linea in bloque:
        for palabra in linea:
            if palabra["x0"] >= x_limite:
                continue
            usadas.append(palabra)
            texto += palabra["text"]
            if "]" in texto:
                return normalizar_sku(texto[1:texto.index("]")]), usadas
    return normalizar_sku(texto.lstrip("[")), usadas


def _cantidad_de_bloque(bloque: list[list[dict]], x_ini: float, x_fin: float) -> tuple[float, str]:
    cantidad, unidad = 0.0, ""
    for linea in bloque:
        for palabra in linea:
            if not (x_ini <= palabra["x0"] < x_fin):
                continue
            if not cantidad and RE_CANTIDAD.match(palabra["text"]):
                cantidad = float(palabra["text"].replace(",", "."))
            elif cantidad and not unidad and palabra["text"].isalpha():
                unidad = palabra["text"].upper()
    return cantidad, unidad


def _anotaciones_aplanadas(palabras: list[dict], pagina: int) -> list[Anotacion]:
    anotaciones: list[Anotacion] = []
    for linea in _agrupar_lineas(palabras):
        texto = " ".join(p["text"] for p in linea)
        centro = sum((p["top"] + p["bottom"]) / 2 for p in linea) / len(linea)
        coincidencia = RE_PEDIDO.match(texto)
        if coincidencia:
            anotaciones.append(_crear_anotacion(texto, coincidencia, pagina, centro))
        elif anotaciones and abs(linea[0]["top"] - anotaciones[-1].centro_y) < 20:
            previa = anotaciones[-1]
            unido = f"{previa.texto} {texto}"
            anotaciones[-1] = _crear_anotacion(unido, RE_PEDIDO.match(unido), pagina, previa.centro_y)
    return anotaciones


def _sin_duplicar(nuevas: list[Anotacion], existentes: list[Anotacion]) -> list[Anotacion]:
    return [
        a for a in nuevas
        if not any(e.pedido == a.pedido and abs(e.centro_y - a.centro_y) < 15 for e in existentes)
    ]


def _renglon_para(centro_y: float, renglones: list[tuple[RenglonPicking, float]]) -> RenglonPicking | None:
    for renglon, limite in renglones:
        if renglon.top - TOLERANCIA_RENGLON <= centro_y < limite:
            return renglon
    return None


# ---------------------------------------------------------------------------
# Reglas
# ---------------------------------------------------------------------------

def _resultado_de(observaciones: list[str]) -> str:
    if any(o.startswith(f"{ERROR}:") for o in observaciones):
        return ERROR
    if any(o.startswith(f"{REVISAR}:") for o in observaciones):
        return REVISAR
    return CORRECTO


class _Verificacion:
    def __init__(self, lineas: list[LineaLiverpool], pickings: list[Picking],
                 referencias: dict[str, vg.Referencia] | None = None,
                 etiquetas: list[ge.EtiquetaEnvio] | None = None):
        self.lineas = lineas
        self.pickings = pickings
        self.referencias = referencias
        self.etiquetas = etiquetas or []
        self.revision_etiquetas: list[RevisionEtiqueta] = []
        self.diferencias_cantidad: list[tuple[Hallazgo, str, str, float, int, list[DetallePedido]]] = []
        self.por_pedido: dict[str, list[LineaLiverpool]] = {}
        for linea in lineas:
            self.por_pedido.setdefault(linea.pedido, []).append(linea)
        self.hallazgos: list[Hallazgo] = []
        self.detalle: list[DetallePedido] = []
        self.advertencias: list[str] = []
        self.cubiertos: set[tuple[str, str]] = set()
        self.pedidos_encontrados: set[str] = set()

    def ejecutar(self) -> Resultado:
        for picking in self.pickings:
            self._advertencias_picking(picking)
        apariciones = self._contar_apariciones()
        for picking in self.pickings:
            for renglon in picking.renglones:
                self._verificar_renglon(renglon, picking, apariciones)
        sin_picking = self._pedidos_sin_picking() if self.pickings else []
        unidades = self._unidades_picking()
        self._revisar_etiquetas(unidades)
        en_revision = self._pedidos_en_revision()
        guias = vg.revisar_guias(self.por_pedido, self.referencias, unidades, en_revision)
        self._hallazgos_guias(guias)
        self._cobertura_etiquetas(guias)
        self._conciliar_cantidades()
        return self._resultado(sin_picking, vg.ordenar(guias), en_revision)

    def _hallazgo(self, nivel: str, picking: str, pedido: str, sku: str, problema: str, detalle: DetallePedido | None = None):
        self.hallazgos.append(Hallazgo(nivel, picking, pedido, sku, problema))
        if detalle is not None:
            detalle.observaciones.append(f"{nivel}: {problema}")
            if nivel == ERROR or detalle.resultado == CORRECTO:
                detalle.resultado = nivel

    def _advertencias_picking(self, picking: Picking):
        if not picking.renglones:
            self.advertencias.append(f"{picking.archivo}: no se encontraron renglones de productos en el PDF.")
        if picking.total_anotaciones == 0:
            self.advertencias.append(f"{picking.archivo}: el PDF no contiene anotaciones de pedidos reconocibles.")
        for anotacion in picking.anotaciones_huerfanas:
            self._hallazgo(
                REVISAR, picking.nombre, anotacion.pedido, "",
                f"La anotación de la página {anotacion.pagina} no está dentro de ningún renglón de producto; asígnala manualmente.",
            )

    def _contar_apariciones(self) -> dict[tuple[str, str], list[str]]:
        apariciones: dict[tuple[str, str], list[str]] = {}
        for picking in self.pickings:
            for renglon in picking.renglones:
                for anotacion in renglon.anotaciones:
                    apariciones.setdefault((anotacion.pedido, renglon.sku), []).append(picking.archivo)
        return apariciones

    def _verificar_renglon(self, renglon: RenglonPicking, picking: Picking, apariciones):
        if not renglon.anotaciones:
            detalle = self._nuevo_detalle(renglon, None, None)
            self._hallazgo(ERROR, picking.nombre, "", renglon.sku,
                           f"Renglón sin número de pedido ({formatear_cantidad(renglon.cantidad)} {renglon.unidad}).".replace(" )", ")"),
                           detalle)
            return

        suma, cantidad_verificable, vistos = 0, True, set()
        detalles = []
        for anotacion in renglon.anotaciones:
            lineas_pedido = self.por_pedido.get(anotacion.pedido)
            lineas_sku = [l for l in (lineas_pedido or []) if l.sku == renglon.sku]
            detalle = self._nuevo_detalle(renglon, anotacion, lineas_sku or lineas_pedido)
            detalles.append(detalle)
            self._regla_duplicados(anotacion, renglon, picking, apariciones, detalle)

            if not lineas_pedido:
                cantidad_verificable = False
                self._regla_existencia(anotacion, renglon, picking, detalle)
                continue

            self.pedidos_encontrados.add(anotacion.pedido)
            if not lineas_sku:
                cantidad_verificable = False
                esperados = ", ".join(sorted({l.sku for l in lineas_pedido}))
                self._hallazgo(ERROR, picking.nombre, anotacion.pedido, renglon.sku,
                               f"El SKU del picking ({renglon.sku}) no coincide con el SKU de Liverpool ({esperados}).", detalle)
            else:
                self.cubiertos.add((anotacion.pedido, renglon.sku))
                cantidad_lv = sum(l.cantidad for l in lineas_sku)
                if (anotacion.pedido, renglon.sku) not in vistos:
                    suma += cantidad_lv
                    vistos.add((anotacion.pedido, renglon.sku))
                if anotacion.piezas is not None and anotacion.piezas != cantidad_lv:
                    self._hallazgo(ERROR, picking.nombre, anotacion.pedido, renglon.sku,
                                   f"La anotación indica {anotacion.piezas} piezas y Liverpool indica {cantidad_lv}.", detalle)
                    self.diferencias_cantidad.append((self.hallazgos[-1], anotacion.pedido, renglon.sku,
                                                      anotacion.piezas, cantidad_lv, [detalle]))

            self._regla_nombre(anotacion, lineas_pedido, picking, renglon, detalle)
            self._regla_estado(lineas_pedido, picking, anotacion.pedido, renglon.sku, detalle)

        self._regla_cantidad(renglon, picking, suma, cantidad_verificable, detalles)

    def _nuevo_detalle(self, renglon, anotacion, lineas) -> DetallePedido:
        lineas = lineas or []
        primera = lineas[0] if lineas else None
        detalle = DetallePedido(
            resultado=CORRECTO,
            picking=renglon.picking,
            pagina=renglon.pagina,
            pedido=anotacion.pedido if anotacion else "",
            sku_pdf=renglon.sku,
            cantidad_renglon=renglon.cantidad,
            anotacion=anotacion.texto if anotacion else "",
            cliente_liverpool=primera.cliente if primera else "",
            sku_liverpool=", ".join(sorted({l.sku for l in lineas})),
            cantidad_liverpool=str(sum(l.cantidad for l in lineas)) if lineas else "",
            envio=primera.envio if primera else "",
            guia=(primera.guia or (f"Sin guía ({primera.estado_guia})" if primera.estado_guia else "")) if primera else "",
            productos_pedido=len(self.por_pedido.get(anotacion.pedido, [])) if anotacion else 0,
        )
        self.detalle.append(detalle)
        return detalle

    def _regla_existencia(self, anotacion: Anotacion, renglon: RenglonPicking, picking: Picking, detalle):
        problema = "El número de pedido no existe en Liverpool."
        sugerencias = self._sugerir_pedido(anotacion, renglon.sku)
        if len(sugerencias) == 1:
            problema += f" ¿Quisiste decir {sugerencias[0]}?"
        elif sugerencias:
            problema += f" Posibles pedidos: {', '.join(sugerencias)}."
        self._hallazgo(ERROR, picking.nombre, anotacion.pedido, renglon.sku, problema, detalle)

    def _sugerir_pedido(self, anotacion: Anotacion, sku: str) -> list[str]:
        candidatos = []
        for pedido, lineas in self.por_pedido.items():
            distancia = digitos_diferentes(anotacion.pedido, pedido)
            mismo_sku = any(l.sku == sku for l in lineas)
            similitud = max(similitud_nombres(anotacion.nombre, l.cliente) for l in lineas) if anotacion.nombre else 0.0
            nombre_parecido = similitud >= SIMILITUD_NOMBRE_SUGERENCIA
            if distancia <= MAX_DIGITOS_DIFERENTES and (mismo_sku or nombre_parecido):
                candidatos.append((distancia - mismo_sku - similitud, pedido))
            elif mismo_sku and similitud >= 0.85:
                candidatos.append((MAX_DIGITOS_DIFERENTES + 1 - similitud, pedido))
        candidatos.sort()
        if len(candidatos) > 1 and candidatos[0][0] < candidatos[1][0] - 0.5:
            return [candidatos[0][1]]
        return [p for _, p in candidatos[:3]]

    def _regla_duplicados(self, anotacion, renglon, picking, apariciones, detalle):
        archivos = apariciones.get((anotacion.pedido, renglon.sku), [])
        if len(archivos) < 2:
            return
        otros = [a for a in archivos if a != picking.archivo]
        if otros:
            problema = f"Pedido duplicado entre varios PDF ({', '.join(sorted(set(archivos)))})."
        else:
            problema = f"Pedido duplicado dentro del mismo PDF ({len(archivos)} veces)."
        self._hallazgo(ERROR, picking.nombre, anotacion.pedido, renglon.sku, problema, detalle)

    def _regla_nombre(self, anotacion, lineas_pedido, picking, renglon, detalle):
        cliente = lineas_pedido[0].cliente
        if not anotacion.nombre:
            self._hallazgo(REVISAR, picking.nombre, anotacion.pedido, renglon.sku,
                           f"La anotación no incluye nombre del cliente. Liverpool: «{cliente}».", detalle)
        elif not nombres_coinciden(anotacion.nombre, cliente):
            self._hallazgo(REVISAR, picking.nombre, anotacion.pedido, renglon.sku,
                           f"El nombre no coincide. Anotación: «{anotacion.nombre}». Liverpool: «{cliente}».", detalle)

    def _regla_estado(self, lineas_pedido, picking, pedido, sku, detalle):
        if any(l.reembolsado for l in lineas_pedido):
            self._hallazgo(ERROR, picking.nombre, pedido, sku, "Pedido reembolsado: no surtir.", detalle)
            return
        estados = sorted({l.estado for l in lineas_pedido if normalizar_texto(l.estado) != normalizar_texto(ESTADO_ESPERADO)})
        if estados:
            self._hallazgo(REVISAR, picking.nombre, pedido, sku,
                           f"Estado en Liverpool: «{', '.join(e or 'vacío' for e in estados)}» (se esperaba «{ESTADO_ESPERADO}»).",
                           detalle)

    def _regla_cantidad(self, renglon, picking, suma, verificable, detalles):
        if not verificable:
            for detalle in detalles:
                detalle.observaciones.append("Cantidad del renglón no verificable: hay pedidos o SKU sin coincidencia.")
            return
        if abs(suma - renglon.cantidad) > 1e-9:
            pedidos = ", ".join(a.pedido for a in renglon.anotaciones)
            problema = (f"La cantidad del renglón ({formatear_cantidad(renglon.cantidad)}) no coincide con la suma "
                        f"de Liverpool ({suma}) para los pedidos anotados.")
            self.hallazgos.append(Hallazgo(ERROR, picking.nombre, pedidos, renglon.sku, problema))
            for detalle in detalles:
                detalle.observaciones.append(f"{ERROR}: {problema}")
                detalle.resultado = ERROR
            if len(renglon.anotaciones) == 1:
                self.diferencias_cantidad.append((self.hallazgos[-1], renglon.anotaciones[0].pedido, renglon.sku,
                                                  renglon.cantidad, suma, detalles))

    def _unidades_picking(self) -> dict[tuple[str, str], str]:
        unidades = {}
        for picking in self.pickings:
            for renglon in picking.renglones:
                for anotacion in renglon.anotaciones:
                    unidades[(anotacion.pedido, renglon.sku)] = renglon.unidad
        return unidades

    def _pedidos_en_revision(self) -> set[str]:
        if not self.pickings and not self.etiquetas:
            return set(self.por_pedido)
        en_etiquetas = {r.pedido for r in self.revision_etiquetas if r.pedido in self.por_pedido}
        return self.pedidos_encontrados | en_etiquetas

    def _hallazgos_guias(self, revisiones: list[vg.RevisionGuias]) -> None:
        for revision in revisiones:
            if not revision.en_revision or revision.resultado != REVISAR:
                continue
            problema = "Datos de paquetes en Liverpool: " + " ".join(vg.problemas(revision))
            skus = ", ".join(sorted({l.sku for l in self.por_pedido[revision.pedido]}))
            self.hallazgos.append(Hallazgo(REVISAR, "Reporte Liverpool", revision.pedido, skus, problema, ORIGEN_EXCEL))
            for detalle in self.detalle:
                if detalle.pedido == revision.pedido:
                    detalle.observaciones.append(f"{REVISAR}: {problema}")
                    if detalle.resultado == CORRECTO:
                        detalle.resultado = REVISAR

    # --- Segunda revisión: guías de envío -----------------------------------

    def _hallazgo_etiqueta(self, revision: RevisionEtiqueta, nivel: str, problema: str, sku: str = ""):
        documento = f"{revision.archivo} p. {revision.pagina}"
        self.hallazgos.append(Hallazgo(nivel, documento, revision.pedido, sku, problema, ORIGEN_GUIAS))
        revision.observaciones.append(f"{nivel}: {problema}")
        if nivel == ERROR or revision.resultado == CORRECTO:
            revision.resultado = nivel

    def _unidades_liverpool(self, contenido: ge.Contenido, linea: LineaLiverpool | None, unidad_picking: str) -> float:
        if contenido.pares is not None:
            return contenido.pares
        if contenido.unidad == "PIEZA" and linea is not None and vg.es_producto_en_par(linea.producto, unidad_picking):
            return contenido.cantidad / 2
        return contenido.cantidad

    def _revisar_etiquetas(self, unidades: dict[tuple[str, str], str]) -> None:
        por_rastreo: dict[str, list[RevisionEtiqueta]] = {}
        pendientes = []
        for etiqueta in self.etiquetas:
            revision = RevisionEtiqueta(
                resultado=CORRECTO, archivo=etiqueta.archivo, pagina=etiqueta.pagina, paqueteria=etiqueta.paqueteria,
                rastreo=etiqueta.rastreo, pedido=etiqueta.pedido,
                identificacion="Referencia de la guía" if etiqueta.pedido else "",
                contenido=etiqueta.texto_contenido, peso=f"{formatear_cantidad(etiqueta.peso)} kg" if etiqueta.peso else "",
                peso_estimado="", medidas="×".join(formatear_cantidad(m) for m in etiqueta.medidas) + " cm" if etiqueta.medidas else "",
            )
            self.revision_etiquetas.append(revision)
            if etiqueta.rastreo:
                por_rastreo.setdefault(etiqueta.rastreo, []).append(revision)
            if etiqueta.pedido:
                self._revisar_etiqueta(etiqueta, revision, unidades)
            else:
                pendientes.append((etiqueta, revision))

        for etiqueta, revision in pendientes:
            self._asignar_por_sku(etiqueta, revision, unidades)

        for rastreo, repetidas in por_rastreo.items():
            if len(repetidas) > 1:
                paginas = ", ".join(f"{r.archivo} p. {r.pagina}" for r in repetidas)
                for revision in repetidas:
                    self._hallazgo_etiqueta(revision, ERROR, f"Guía duplicada: el rastreo {rastreo} aparece en {paginas}.")

    def _asignar_por_sku(self, etiqueta: ge.EtiquetaEnvio, revision: RevisionEtiqueta, unidades) -> None:
        if not etiqueta.contenido:
            revision.identificacion = "No identificado"
            self._hallazgo_etiqueta(revision, REVISAR, "La página no tiene referencia de pedido ni nota de contenido legibles; revísala manualmente.")
            return
        skus = {c.sku for c in etiqueta.contenido}
        base = self.pedidos_encontrados if self.pickings else set(self.por_pedido)
        con_guia = {r.pedido for r in self.revision_etiquetas if r.pedido}
        candidatos = sorted(p for p in base if skus <= {l.sku for l in self.por_pedido.get(p, [])})
        sin_guia = [p for p in candidatos if p not in con_guia]
        elegidos = sin_guia or candidatos
        if len(elegidos) == 1:
            etiqueta.pedido = revision.pedido = elegidos[0]
            revision.identificacion = "Por SKU (la guía no tiene referencia legible)"
            revision.observaciones.append("Pedido identificado por el SKU de la nota; confirma visualmente la guía.")
            self._revisar_etiqueta(etiqueta, revision, unidades)
        else:
            revision.identificacion = "No identificado"
            detalle = f" Posibles pedidos: {', '.join(elegidos[:5])}." if elegidos else ""
            self._hallazgo_etiqueta(revision, REVISAR,
                                    f"No se pudo identificar el pedido de la guía ({etiqueta.texto_contenido}).{detalle}",
                                    ", ".join(sorted(skus)))

    def _revisar_etiqueta(self, etiqueta: ge.EtiquetaEnvio, revision: RevisionEtiqueta, unidades) -> None:
        lineas = self.por_pedido.get(etiqueta.pedido)
        sku_principal = etiqueta.contenido[0].sku if etiqueta.contenido else ""
        if not lineas:
            sugerencias = self._sugerir_pedido(Anotacion("", etiqueta.pedido, "", None, 0, 0), sku_principal)
            extra = f" ¿Quisiste decir {sugerencias[0]}?" if len(sugerencias) == 1 else (
                f" Posibles pedidos: {', '.join(sugerencias)}." if sugerencias else "")
            self._hallazgo_etiqueta(revision, ERROR, "La guía es de un pedido que no existe en Liverpool." + extra, sku_principal)
            return
        if self.pickings and etiqueta.pedido not in self.pedidos_encontrados:
            self._hallazgo_etiqueta(revision, REVISAR, "La guía es de un pedido que no aparece en los picking cargados.", sku_principal)
        if any(l.reembolsado for l in lineas):
            self._hallazgo_etiqueta(revision, ERROR, "Pedido reembolsado: no enviar.", sku_principal)
        if not etiqueta.contenido:
            revision.observaciones.append("La guía no tiene nota de contenido legible.")
            return

        por_sku = {l.sku: l for l in lineas}
        peso_estimado, peso_completo, total_unidades = 0.0, True, 0.0
        for contenido in etiqueta.contenido:
            linea = por_sku.get(contenido.sku)
            if linea is None:
                self._hallazgo_etiqueta(revision, ERROR,
                                        f"La guía indica {contenido.sku}, que no pertenece al pedido "
                                        f"(SKU en Liverpool: {', '.join(sorted(por_sku))}).", contenido.sku)
                peso_completo = False
                continue
            unidad = unidades.get((etiqueta.pedido, contenido.sku), "")
            cantidad = self._unidades_liverpool(contenido, linea, unidad)
            total_unidades += cantidad
            ref = (self.referencias or {}).get(contenido.sku)
            peso = vg.peso_unidad(contenido.sku, linea.producto, unidad, ref)
            if peso is None:
                peso_completo = False
            else:
                peso_estimado += peso * cantidad
                if vg.es_producto_en_par(linea.producto, unidad) and cantidad >= 1 and contenido.unidad != "PIEZA" \
                        and peso / 2 >= vg.PESO_UNIDAD_GUIA_INDIVIDUAL:
                    self._hallazgo_etiqueta(revision, REVISAR,
                                            f"La guía lleva un par de {contenido.sku} ({formatear_cantidad(peso / 2)} kg por pieza): "
                                            f"con {formatear_cantidad(vg.PESO_UNIDAD_GUIA_INDIVIDUAL)} kg o más va una guía por pieza.",
                                            contenido.sku)
            if len(etiqueta.contenido) == 1 and cantidad == 1 and ref and not ref.ambigua and len(ref.cajas) == 1 \
                    and etiqueta.medidas and ref.cajas[0].medidas:
                if not all(abs(a - b) <= max(vg.TOLERANCIA_MEDIDA_CM, b * vg.TOLERANCIA_RELATIVA)
                           for a, b in zip(sorted(etiqueta.medidas), ref.cajas[0].medidas)):
                    self._hallazgo_etiqueta(revision, REVISAR,
                                            f"Medidas de la guía {revision.medidas}; la tabla indica "
                                            f"{ref.cajas[0].texto().split(' cm')[0]} cm.", contenido.sku)

        if peso_completo and peso_estimado:
            revision.peso_estimado = f"{formatear_cantidad(round(peso_estimado, 1))} kg"
            if etiqueta.peso is not None and abs(etiqueta.peso - peso_estimado) > max(
                    vg.TOLERANCIA_PESO_KG, peso_estimado * vg.TOLERANCIA_RELATIVA):
                self._hallazgo_etiqueta(revision, REVISAR,
                                        f"Peso de la guía {revision.peso}; el contenido pesa ≈ {revision.peso_estimado}.",
                                        sku_principal)
        if etiqueta.peso is not None and etiqueta.peso > vg.PESO_MAXIMO_GUIA and total_unidades > 1:
            self._hallazgo_etiqueta(revision, REVISAR,
                                    f"La guía agrupa varios productos y pesa {revision.peso}: rebasa "
                                    f"{formatear_cantidad(vg.PESO_MAXIMO_GUIA)} kg.", sku_principal)

    def _cobertura_etiquetas(self, revisiones: list[vg.RevisionGuias]) -> None:
        if not self.etiquetas:
            return
        por_pedido: dict[str, list[ge.EtiquetaEnvio]] = {}
        for etiqueta in self.etiquetas:
            if etiqueta.pedido in self.por_pedido:
                por_pedido.setdefault(etiqueta.pedido, []).append(etiqueta)
        unidades = self._unidades_picking()
        for revision in revisiones:
            etiquetas = por_pedido.get(revision.pedido, [])
            revision.guias_generadas = str(len(etiquetas)) if etiquetas else ""
            if not etiquetas:
                if self.pickings and revision.pedido in self.pedidos_encontrados:
                    self._hallazgo_pedido_guias(revision, "El pedido está en el picking pero no tiene guía de envío entre los documentos.")
                continue
            if len(etiquetas) < revision.minimo:
                self._hallazgo_pedido_guias(revision, f"Se generaron {len(etiquetas)} guías; se necesitan al menos {revision.minimo}.")
            elif revision.completa and revision.maximo and len(etiquetas) > revision.maximo:
                self._hallazgo_pedido_guias(revision, f"Se generaron {len(etiquetas)} guías; bastan {revision.maximo}.")
            else:
                revision.observaciones.append(f"Las {len(etiquetas)} guías generadas cumplen la estimación ({revision.guias_estimadas}).")
            guias_lv = vg._a_float(revision.guias_liverpool)
            if guias_lv and int(guias_lv) != len(etiquetas):
                revision.observaciones.append(f"Liverpool registra {int(guias_lv)} guías y se generaron {len(etiquetas)}.")
            if all(e.contenido for e in etiquetas):
                lineas = {l.sku: l for l in self.por_pedido[revision.pedido]}
                cubiertas: dict[str, float] = {}
                for etiqueta in etiquetas:
                    for contenido in etiqueta.contenido:
                        cubiertas[contenido.sku] = cubiertas.get(contenido.sku, 0) + self._unidades_liverpool(
                            contenido, lineas.get(contenido.sku), unidades.get((revision.pedido, contenido.sku), ""))
                for sku, linea in lineas.items():
                    cantidad_lv = sum(l.cantidad for l in self.por_pedido[revision.pedido] if l.sku == sku)
                    enviado = cubiertas.get(sku, 0)
                    if enviado < cantidad_lv:
                        self._hallazgo_pedido_guias(revision, f"Las guías cubren {formatear_cantidad(enviado)} de {cantidad_lv} unidades de {sku}.", sku)
                    elif enviado > cantidad_lv:
                        self._hallazgo_pedido_guias(revision, f"Las guías indican {formatear_cantidad(enviado)} unidades de {sku} y Liverpool {cantidad_lv}.", sku, ERROR)
            else:
                revision.observaciones.append("Alguna guía del pedido no tiene nota de contenido; no se pudo comprobar qué productos cubren.")

    def _conciliar_cantidades(self) -> None:
        """Si la guía de envío aclara que las piezas del picking equivalen a los pares de Liverpool
        (por ejemplo «2 Piezas (1 Par)»), la diferencia de cantidad baja de ERROR a REVISAR."""
        for hallazgo, pedido, sku, cantidad_pdf, cantidad_lv, detalles in self.diferencias_cantidad:
            contenidos = [c for e in self.etiquetas if e.pedido == pedido for c in e.contenido if c.sku == sku]
            if not contenidos or any(c.pares is None for c in contenidos):
                continue
            piezas = sum(c.cantidad for c in contenidos)
            pares = sum(c.pares for c in contenidos)
            if piezas != cantidad_pdf or pares != cantidad_lv:
                continue
            original = hallazgo.problema
            hallazgo.nivel = REVISAR
            hallazgo.problema = (f"{original} La guía de envío indica {piezas} piezas ({pares} par): el picking cuenta "
                                 "piezas y Liverpool pares. Confirma la unidad de venta.")
            for detalle in detalles:
                detalle.observaciones = [f"{REVISAR}: {hallazgo.problema}" if o == f"{ERROR}: {original}" else o
                                         for o in detalle.observaciones]
                detalle.resultado = _resultado_de(detalle.observaciones)

    def _hallazgo_pedido_guias(self, revision: vg.RevisionGuias, problema: str, sku: str = "", nivel: str = REVISAR) -> None:
        skus = sku or ", ".join(sorted({l.sku for l in self.por_pedido[revision.pedido]}))
        self.hallazgos.append(Hallazgo(nivel, "Guías de envío", revision.pedido, skus, problema, ORIGEN_GUIAS))
        revision.observaciones.append(f"{nivel}: {problema}")
        revision.resultado = REVISAR if revision.resultado == CORRECTO else revision.resultado

    def _pedidos_sin_picking(self) -> list[SinPicking]:
        faltantes = []
        for linea in self.lineas:
            if (linea.pedido, linea.sku) in self.cubiertos:
                continue
            if linea.pedido in self.pedidos_encontrados:
                observacion = "El pedido aparece en un picking, pero falta este SKU."
            else:
                observacion = "El pedido no aparece en ningún picking."
            if linea.reembolsado:
                observacion += " Pedido reembolsado."
            faltantes.append(SinPicking(linea.pedido, linea.sku, linea.cantidad, linea.cliente, linea.producto,
                                        linea.estado, linea.fecha_limite, linea.envio, observacion))
        faltantes.sort(key=lambda s: (parsear_fecha_limite(s.fecha_limite), s.pedido))
        return faltantes

    def _resultado(self, sin_picking: list[SinPicking], guias: list[vg.RevisionGuias], en_revision: set[str]) -> Resultado:
        con_hallazgos = {p.strip() for h in self.hallazgos for p in h.pedido.split(",") if p.strip()}
        revisados = {p for p in en_revision if p in self.por_pedido}
        orden_origen = {ORIGEN_EXCEL: 0, ORIGEN_PICKING: 1, ORIGEN_GUIAS: 2}
        return Resultado(
            pedidos_liverpool=len(self.por_pedido),
            verificados=len(revisados),
            correctos=sum(1 for p in revisados if p not in con_hallazgos),
            errores=sum(1 for h in self.hallazgos if h.nivel == ERROR),
            revisar=sum(1 for h in self.hallazgos if h.nivel == REVISAR),
            sin_picking=len({s.pedido for s in sin_picking}) if self.pickings else None,
            hallazgos=sorted(self.hallazgos, key=lambda h: (h.nivel != ERROR, orden_origen.get(h.origen, 3), h.picking, h.pedido)),
            detalle=self.detalle,
            pedidos_sin_picking=sin_picking,
            advertencias=self.advertencias,
            pickings=[p.nombre for p in self.pickings],
            guias=guias,
            guias_por_revisar=sum(1 for g in guias if g.en_revision and g.resultado == REVISAR),
            picking_cargado=bool(self.pickings),
            etiquetas=self.revision_etiquetas,
            guias_envio_cargadas=bool(self.etiquetas),
        )


def verificar(excel: Fuente, pdfs: Iterable = (), medidas: Fuente | None = None, guias: Iterable = ()) -> Resultado:
    """Primera revisión: Excel de Liverpool contra la tabla de medidas (opcional).
    Segunda revisión: PDF de picking de Odoo y PDF de guías de envío (ambos opcionales)."""
    referencias = leer_medidas(medidas) if medidas is not None else None
    lineas = leer_liverpool(excel)
    if not lineas:
        raise DocumentoInvalido("El Excel de Liverpool no contiene pedidos.")
    pickings = [leer_picking(f, n) for n, f in (_con_nombre(pdf, i) for i, pdf in enumerate(pdfs))]
    etiquetas = []
    for nombre, fuente in (_con_nombre(pdf, i) for i, pdf in enumerate(guias)):
        etiquetas.extend(leer_guias(fuente, nombre))
    resultado = _Verificacion(lineas, pickings, referencias, etiquetas).ejecutar()
    resultado.medidas_cargadas = referencias is not None
    return resultado


def leer_guias(pdf: Fuente, nombre_archivo: str) -> list[ge.EtiquetaEnvio]:
    datos = _leer_bytes(pdf)
    try:
        return ge.leer_etiquetas(datos, nombre_archivo)
    except Exception as exc:
        raise DocumentoInvalido(f"El documento PDF «{nombre_archivo}» está dañado o no puede leerse.") from exc


def _leer_bytes(fuente: Fuente) -> bytes:
    manejador = _abrir_binario(fuente)
    try:
        return manejador.read()
    finally:
        if isinstance(fuente, (str, Path)):
            manejador.close()


def clasificar(nombre: str, fuente: Fuente) -> str:
    """Identifica el tipo de documento por su contenido (no por su nombre)."""
    if nombre.lower().endswith(".xlsx"):
        try:
            leer_liverpool(fuente)
            return TIPO_LIVERPOOL
        except DocumentoInvalido as exc:
            if "no parece ser el reporte" not in str(exc):
                raise
        try:
            leer_medidas(fuente)
            return TIPO_MEDIDAS
        except DocumentoInvalido:
            return TIPO_DESCONOCIDO
    if nombre.lower().endswith(".pdf"):
        datos = _leer_bytes(fuente)
        try:
            with pdfplumber.open(io.BytesIO(datos)) as documento:
                texto = "\n".join((pagina.extract_text() or "") for pagina in documento.pages[:3])
        except Exception as exc:
            raise DocumentoInvalido(f"El documento PDF «{nombre}» está dañado o no puede leerse.") from exc
        if RE_PICKING.search(texto) or ("CANTIDAD" in texto.upper() and re.search(r"\[[A-Z]{1,4}-", texto)):
            return TIPO_PICKING
        if ge.parece_guia(texto):
            return TIPO_GUIAS
    return TIPO_DESCONOCIDO


def analizar(documentos: Iterable[tuple[str, Fuente]]) -> Resultado:
    """Recibe los documentos en cualquier orden y cantidad, los identifica y ejecuta las revisiones posibles.
    El único obligatorio es el reporte de pedidos de Liverpool."""
    grupos: dict[str, list[tuple[str, Fuente]]] = {}
    for nombre, fuente in documentos:
        grupos.setdefault(clasificar(nombre, fuente), []).append((nombre, fuente))
    liverpool = grupos.get(TIPO_LIVERPOOL, [])
    medidas = grupos.get(TIPO_MEDIDAS, [])
    if not liverpool:
        excels = [n for n, _ in grupos.get(TIPO_DESCONOCIDO, []) if n.lower().endswith(".xlsx")]
        if excels:
            raise DocumentoInvalido("El Excel no parece ser el reporte de pedidos de Liverpool. "
                                    f"Revisa: {', '.join(excels)}.")
        raise DocumentoInvalido("No se encontró el reporte de pedidos de Liverpool entre los documentos. "
                                "Es el documento principal y es obligatorio.")
    if len(liverpool) > 1:
        raise DocumentoInvalido("Se encontraron varios reportes de Liverpool "
                                f"({', '.join(n for n, _ in liverpool)}). Sube solo uno por verificación.")
    if len(medidas) > 1:
        raise DocumentoInvalido(f"Se encontraron varias tablas de medidas ({', '.join(n for n, _ in medidas)}). Sube solo una.")

    resultado = verificar(
        liverpool[0][1],
        grupos.get(TIPO_PICKING, []),
        medidas[0][1] if medidas else None,
        grupos.get(TIPO_GUIAS, []),
    )
    resultado.documentos = _resumen_documentos(grupos, resultado)
    for nombre, _ in grupos.get(TIPO_DESCONOCIDO, []):
        resultado.advertencias.append(f"«{nombre}» no se reconoció como reporte de Liverpool, tabla de medidas, "
                                      "picking ni guías de envío; no se usó.")
    return resultado


def _resumen_documentos(grupos, resultado: Resultado) -> list[DocumentoAnalizado]:
    resumen = []
    for tipo in (TIPO_LIVERPOOL, TIPO_MEDIDAS, TIPO_PICKING, TIPO_GUIAS, TIPO_DESCONOCIDO):
        for nombre, _ in grupos.get(tipo, []):
            if tipo == TIPO_LIVERPOOL:
                detalle = f"{resultado.pedidos_liverpool} pedidos · primera revisión"
            elif tipo == TIPO_MEDIDAS:
                detalle = "Referencia de guías, pesos y medidas · primera revisión"
            elif tipo == TIPO_PICKING:
                detalle = "Segunda revisión"
            elif tipo == TIPO_GUIAS:
                paginas = sum(1 for e in resultado.etiquetas if e.archivo == nombre)
                detalle = f"{paginas} guías · segunda revisión"
            else:
                detalle = "No se usó"
            resumen.append(DocumentoAnalizado(nombre, tipo, detalle))
    return resumen


def leer_medidas(medidas: Fuente) -> dict[str, vg.Referencia]:
    manejador = _abrir_binario(medidas)
    try:
        referencias = vg.leer_medidas(manejador)
    except Exception as exc:
        raise DocumentoInvalido("La tabla de medidas de equipos no tiene un formato válido.") from exc
    finally:
        if isinstance(medidas, (str, Path)):
            manejador.close()
    if not referencias:
        raise DocumentoInvalido("La tabla de medidas de equipos no contiene equipos reconocibles.")
    return referencias


# ---------------------------------------------------------------------------
# Reporte Excel
# ---------------------------------------------------------------------------

_RELLENOS = {
    ERROR: PatternFill("solid", fgColor="F8D7DA"),
    REVISAR: PatternFill("solid", fgColor="FFF3CD"),
    CORRECTO: PatternFill("solid", fgColor="E3F2D0"),
}
_ENCABEZADO = PatternFill("solid", fgColor="393939")
_PELIGROSOS = ("=", "+", "-", "@", "\t", "\r")


def texto_seguro(valor: object) -> object:
    """Evita que un texto proveniente de los documentos se interprete como fórmula."""
    if isinstance(valor, str) and valor.startswith(_PELIGROSOS):
        return "'" + valor
    return valor


def _escribir_tabla(hoja, encabezados: list[str], filas: list[list], nivel_col: int | None = None, anchos=None):
    hoja.append(encabezados)
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = _ENCABEZADO
        celda.alignment = Alignment(vertical="center", wrap_text=True)
    for fila in filas:
        hoja.append([texto_seguro(v) for v in fila])
        if nivel_col is not None:
            relleno = _RELLENOS.get(fila[nivel_col])
            if relleno:
                for celda in hoja[hoja.max_row]:
                    celda.fill = relleno
    for i, ancho in enumerate(anchos or [], 1):
        hoja.column_dimensions[get_column_letter(i)].width = ancho
    for fila in hoja.iter_rows(min_row=2):
        for celda in fila:
            if isinstance(celda.value, str):
                celda.data_type = "s"
            celda.alignment = Alignment(vertical="top", wrap_text=True)
    hoja.freeze_panes = "A2"
    if filas:
        hoja.auto_filter.ref = hoja.dimensions


def escribir_reporte(resultado: Resultado, destino: Union[str, Path, BinaryIO], generado: datetime | None = None) -> None:
    """Escribe el reporte de verificación en un archivo .xlsx (ruta o flujo binario)."""
    libro = Workbook()
    resumen = libro.active
    resumen.title = "Resumen"
    momento = (generado or datetime.now()).strftime("%d/%m/%Y %H:%M")
    if resultado.errores:
        veredicto = f"Se encontraron {resultado.errores} errores"
    elif resultado.revisar or resultado.advertencias:
        veredicto = "Sin errores, pero hay elementos para revisar"
    else:
        veredicto = "Todo cuadra"
    datos = [
        ("Verificación de pedidos Liverpool – Gymsport", ""),
        ("Generado", momento),
        ("Documentos", "; ".join(f"{d.nombre} ({d.tipo})" for d in resultado.documentos) or ", ".join(resultado.pickings)),
        ("Resultado", veredicto),
        ("", ""),
        ("Pedidos en Liverpool", resultado.pedidos_liverpool),
        ("Verificados", resultado.verificados),
        ("Correctos", resultado.correctos),
        ("Errores", resultado.errores),
        ("Para revisar", resultado.revisar),
        ("Sin picking", resultado.sin_picking if resultado.sin_picking is not None else "no se cargó picking"),
        ("Guías por revisar", resultado.guias_por_revisar),
        ("Tabla de medidas", "cargada" if resultado.medidas_cargadas else "no cargada (solo reglas de empaque)"),
    ]
    for fila in datos:
        resumen.append([texto_seguro(v) for v in fila])
    resumen["A1"].font = Font(bold=True, size=14)
    for fila in range(6, 13):
        resumen.cell(fila, 1).font = Font(bold=True)
    resumen.column_dimensions["A"].width = 28
    resumen.column_dimensions["B"].width = 60
    if resultado.advertencias:
        resumen.append([])
        resumen.append(["Advertencias"])
        resumen.cell(resumen.max_row, 1).font = Font(bold=True)
        for advertencia in resultado.advertencias:
            resumen.append(["", texto_seguro(advertencia)])

    _escribir_tabla(
        libro.create_sheet("Hallazgos"),
        ["Nivel", "Revisión", "Documento", "Pedido", "SKU", "Problema"],
        [[h.nivel, h.origen, h.picking, h.pedido, h.sku, h.problema] for h in resultado.hallazgos],
        nivel_col=0, anchos=[11, 16, 24, 24, 18, 90],
    )
    _escribir_tabla(
        libro.create_sheet("Detalle por pedido"),
        ["Resultado", "Picking", "Página", "Pedido", "SKU PDF", "Cant. renglón", "Anotación", "Cliente Liverpool",
         "SKU Liverpool", "Cant. Liverpool", "Envío", "Guía", "Observaciones"],
        [[d.resultado, d.picking, d.pagina, d.pedido, d.sku_pdf, d.cantidad_renglon, d.anotacion, d.cliente_liverpool,
          d.sku_liverpool, d.cantidad_liverpool, d.envio, d.guia, "\n".join(d.observaciones)] for d in resultado.detalle],
        nivel_col=0, anchos=[11, 18, 8, 13, 16, 9, 40, 30, 16, 9, 16, 18, 70],
    )
    _escribir_tabla(
        libro.create_sheet("Sin picking"),
        ["Fecha límite de envío", "Pedido", "SKU", "Cantidad", "Cliente", "Producto", "Estado", "Envío", "Observación"],
        [[s.fecha_limite, s.pedido, s.sku, s.cantidad, s.cliente, s.producto, s.estado, s.envio, s.observacion]
         for s in resultado.pedidos_sin_picking],
        anchos=[22, 13, 16, 9, 30, 45, 24, 16, 45],
    )
    _escribir_tabla(
        libro.create_sheet("Guías y medidas"),
        ["Resultado", "Pedido", "En revisión", "Productos", "Unidades", "Guías Liverpool", "Guías estimadas",
         "Guías generadas", "Paquetes en Liverpool", "Referencia", "Observaciones"],
        [[g.resultado, g.pedido, "Sí" if g.en_revision else "No", g.productos, g.unidades, g.guias_liverpool,
          g.guias_estimadas, g.guias_generadas, g.paquetes_liverpool, g.referencia, "\n".join(g.observaciones)]
         for g in resultado.guias],
        nivel_col=0, anchos=[11, 13, 10, 34, 9, 10, 12, 10, 45, 30, 80],
    )
    if resultado.guias_envio_cargadas:
        _escribir_tabla(
            libro.create_sheet("Guías de envío"),
            ["Resultado", "Archivo", "Página", "Paquetería", "Rastreo", "Pedido", "Identificación", "Contenido",
             "Peso", "Peso estimado", "Medidas", "Observaciones"],
            [[e.resultado, e.archivo, e.pagina, e.paqueteria, e.rastreo, e.pedido, e.identificacion, e.contenido,
              e.peso, e.peso_estimado, e.medidas, "\n".join(e.observaciones)] for e in resultado.etiquetas],
            nivel_col=0, anchos=[11, 28, 8, 12, 18, 13, 22, 34, 9, 10, 14, 70],
        )
    libro.save(destino)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Verifica pedidos de Liverpool contra picking de Odoo y guías de envío.")
    parser.add_argument("documentos", nargs="+", help="reporte de Liverpool y, opcionalmente, medidas, picking y guías")
    parser.add_argument("-o", "--salida", default="verificacion_pedidos.xlsx")
    args = parser.parse_args(argv)
    try:
        resultado = analizar([(Path(d).name, d) for d in args.documentos])
    except DocumentoInvalido as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 2
    escribir_reporte(resultado, args.salida)
    for documento in resultado.documentos:
        print(f"{documento.tipo}: {documento.nombre}")
    print(f"Pedidos en Liverpool: {resultado.pedidos_liverpool}")
    print(f"Verificados: {resultado.verificados}  Correctos: {resultado.correctos}")
    print(f"Errores: {resultado.errores}  Para revisar: {resultado.revisar}  Sin picking: {resultado.sin_picking}")
    print(f"Guías por revisar: {resultado.guias_por_revisar}")
    for advertencia in resultado.advertencias:
        print(f"Advertencia: {advertencia}")
    print(f"Reporte: {args.salida}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
