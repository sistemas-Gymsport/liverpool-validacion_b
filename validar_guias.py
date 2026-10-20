"""Validación de guías, pesos y medidas de los paquetes que registra Liverpool.

Los paquetes del Excel de Liverpool se tratan como datos A VALIDAR, no como referencia:
cuando Liverpool no tiene medidas guardadas exporta valores genéricos (10 x 10 x 10 cm y 1 kg),
ceros o listas vacías. La referencia es la tabla interna de medidas de equipos (opcional) más
las reglas de empaque de Gymsport:

- Productos de un mismo pedido pueden compartir guía mientras no rebasen PESO_MAXIMO_GUIA kg.
- Los metros de piso van de dos en dos por guía.
- Las mancuernas/discos vendidos en par con PESO_UNIDAD_GUIA_INDIVIDUAL kg o más por unidad
  llevan una guía por unidad.
"""

from __future__ import annotations

import io
import math
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import BinaryIO, Iterable, Union

from openpyxl import load_workbook

PESO_MAXIMO_GUIA = 30.0
PESO_UNIDAD_GUIA_INDIVIDUAL = 17.0
LIBRA_KG = 0.4536
TOLERANCIA_PESO_KG = 2.0
TOLERANCIA_MEDIDA_CM = 5.0
TOLERANCIA_RELATIVA = 0.15
TOLERANCIA_RELATIVA_TOTAL = 0.10
PAQUETE_GENERICO = (10.0, 10.0, 10.0, 1.0)
NOTAS_INFORMATIVAS = ("Sin medidas de referencia", "Referencia ambigua", "No se cargó", "Una guía por unidad")

RE_SKU_REFERENCIA = re.compile(r"^[A-Z]{1,4}-[A-Z0-9./-]+$")
RE_PESO_NOMBRE = re.compile(r"(\d+(?:[.,]\d+)?)\s*(kg|kgs|k|lb|lbs|libras?)\b", re.I)
RE_NUMERO_SKU = re.compile(r"(\d+(?:\.\d+)?)(?!.*\d)")

Fuente = Union[str, Path, bytes, BinaryIO]


@dataclass(frozen=True)
class Caja:
    peso: float | None
    largo: float | None
    alto: float | None
    ancho: float | None

    @property
    def medidas(self) -> tuple[float, ...] | None:
        valores = (self.largo, self.alto, self.ancho)
        return tuple(sorted(valores)) if all(v is not None for v in valores) else None

    def texto(self) -> str:
        medidas = "×".join(_num(v) for v in (self.largo, self.alto, self.ancho)) if self.medidas else "sin medidas"
        peso = f"{_num(self.peso)} kg" if self.peso is not None else "sin peso"
        return f"{medidas} cm, {peso}"


@dataclass
class Referencia:
    sku: str
    guias: int | None
    cajas: list[Caja]
    ambigua: bool = False
    nota: str = ""


@dataclass
class Paquetes:
    guias: int | None
    cajas: list[Caja]
    texto_guias: str


@dataclass
class RevisionGuias:
    pedido: str
    productos: str
    unidades: int
    en_revision: bool
    guias_liverpool: str
    guias_estimadas: str
    paquetes_liverpool: str
    referencia: str
    resultado: str
    observaciones: list[str] = field(default_factory=list)
    guias_generadas: str = ""
    minimo: int = 0
    maximo: int = 0
    completa: bool = True


def _num(valor: float | None) -> str:
    if valor is None:
        return "?"
    return str(int(valor)) if float(valor).is_integer() else f"{valor:.1f}"


def _a_float(valor) -> float | None:
    if valor is None or valor == "":
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = str(valor).strip().replace(",", ".")
    try:
        return float(texto)
    except ValueError:
        return None


# ---------------------------------------------------------------------------
# Tabla de medidas de equipos
# ---------------------------------------------------------------------------

def leer_medidas(fuente: Fuente) -> dict[str, Referencia]:
    """Lee la tabla de medidas. Admite varias hojas y varias tablas lado a lado (bloques de 6 columnas:
    equipo, guías, peso kg, largo, alto, ancho)."""
    if isinstance(fuente, (bytes, bytearray)):
        fuente = io.BytesIO(fuente)
    elif hasattr(fuente, "seek"):
        fuente.seek(0)
    libro = load_workbook(fuente, read_only=True, data_only=True)
    filas_por_sku: dict[str, list[tuple]] = {}
    try:
        for hoja in libro.worksheets:
            hoja.reset_dimensions()
            for fila in hoja.iter_rows(values_only=True):
                fila = list(fila)
                for inicio in range(len(fila)):
                    sku = fila[inicio]
                    if not isinstance(sku, str):
                        continue
                    sku = re.sub(r"\s+", "", sku).upper()
                    bloque = fila[inicio + 1:inicio + 6] + [None] * 5
                    numeros = [_a_float(v) for v in bloque[:5]]
                    if not RE_SKU_REFERENCIA.match(sku) or not any(n is not None for n in numeros):
                        continue
                    if any(isinstance(v, str) and _a_float(v) is None for v in bloque[:5] if v not in (None, "")):
                        continue
                    filas_por_sku.setdefault(sku, []).append(tuple(numeros))
    finally:
        libro.close()
    return {sku: _consolidar(sku, filas) for sku, filas in filas_por_sku.items()}


def _consolidar(sku: str, filas: list[tuple]) -> Referencia:
    unicas = list(dict.fromkeys(filas))
    guias = {int(f[0]) for f in unicas if f[0] is not None and f[0] >= 1}
    cajas = [Caja(f[1], f[2], f[3], f[4]) for f in unicas]
    if len(unicas) == 1:
        return Referencia(sku, next(iter(guias), None), cajas)
    if len(guias) == 1 and next(iter(guias)) >= len(unicas) and next(iter(guias)) > 1:
        return Referencia(sku, next(iter(guias)), cajas, nota="equipo en varias cajas")
    return Referencia(sku, max(guias) if len(guias) == 1 else None, cajas, ambigua=True,
                      nota=f"la tabla de medidas tiene {len(unicas)} versiones distintas")


# ---------------------------------------------------------------------------
# Paquetes de Liverpool
# ---------------------------------------------------------------------------

def _lista(texto: str) -> list[float | None]:
    limpio = (texto or "").strip().strip("[]")
    if not limpio:
        return []
    return [_a_float(v) for v in limpio.split(",")]


def leer_paquetes(largo: str, alto: str, ancho: str, peso: str, guias: str) -> Paquetes:
    columnas = [_lista(largo), _lista(alto), _lista(ancho), _lista(peso)]
    total = max(len(c) for c in columnas)
    cajas = [Caja(*(c[i] if i < len(c) else None for c in (columnas[3], columnas[0], columnas[1], columnas[2])))
             for i in range(total)]
    numero = _a_float(guias)
    return Paquetes(int(numero) if numero is not None else None, cajas, (guias or "").strip())


def es_generico(caja: Caja) -> bool:
    valores = (caja.largo, caja.alto, caja.ancho, caja.peso)
    if all(v in (None, 0) for v in valores):
        return True
    return valores == PAQUETE_GENERICO


# ---------------------------------------------------------------------------
# Estimación de guías
# ---------------------------------------------------------------------------

@dataclass
class _Estimacion:
    fijas: int = 0
    piso: int = 0
    consolidables: list[float] = field(default_factory=list)
    desconocidas: int = 0
    peso_total: float | None = 0.0
    individuales: dict[str, int] = field(default_factory=dict)

    def sumar_peso(self, peso: float | None):
        self.peso_total = None if self.peso_total is None or peso is None else self.peso_total + peso

    @property
    def minimo(self) -> int:
        return max(self.fijas + math.ceil(self.piso / 2) + _empacar(self.consolidables), 1 if self.desconocidas else 0)

    @property
    def maximo(self) -> int:
        return self.fijas + self.piso + len(self.consolidables) + self.desconocidas


def _empacar(pesos: list[float]) -> int:
    contenedores: list[float] = []
    for peso in sorted(pesos, reverse=True):
        for i, ocupado in enumerate(contenedores):
            if ocupado + peso <= PESO_MAXIMO_GUIA:
                contenedores[i] += peso
                break
        else:
            contenedores.append(peso)
    return len(contenedores)


def peso_por_unidad(nombre: str, sku: str) -> float | None:
    coincidencia = RE_PESO_NOMBRE.search(nombre or "")
    if coincidencia:
        valor = float(coincidencia.group(1).replace(",", "."))
        return valor * LIBRA_KG if coincidencia.group(2).lower().startswith(("lb", "libra")) else valor
    numero = RE_NUMERO_SKU.search(sku.split("-", 1)[-1])
    return float(numero.group(1)) if numero else None


def es_par_de_pesas(nombre: str, unidad_picking: str = "") -> bool:
    texto = (nombre or "").lower()
    es_pesa = any(p in texto for p in ("mancuerna", "disco"))
    en_par = unidad_picking.upper() == "PAR" or bool(re.search(r"\bpar\b|\(par\)|2 piezas", texto))
    return es_pesa and en_par


def es_producto_en_par(nombre: str, unidad_picking: str = "") -> bool:
    return unidad_picking.upper() == "PAR" or es_par_de_pesas(nombre, unidad_picking)


def peso_unidad(sku: str, nombre: str, unidad_picking: str, ref: Referencia | None) -> float | None:
    """Peso de una unidad vendida por Liverpool (un par, en productos que se venden en par)."""
    if ref is not None and not ref.ambigua and ref.guias:
        pesos = [c.peso for c in ref.cajas]
        if all(p is not None for p in pesos):
            if ref.guias > 1:
                return sum(pesos) if len(pesos) == ref.guias else pesos[0] * ref.guias
            return pesos[0]
    por_pieza = peso_por_unidad(nombre, sku)
    if por_pieza is not None and es_par_de_pesas(nombre, unidad_picking):
        return por_pieza * 2
    return None


def es_piso(nombre: str, sku: str) -> bool:
    return "PISO" in sku.upper() or bool(re.search(r"\bpiso\b", (nombre or "").lower()))


def _estimar_unidad(est: _Estimacion, sku: str, nombre: str, unidad: str, ref: Referencia | None):
    if es_piso(nombre, sku):
        est.piso += 1
        est.sumar_peso(ref.cajas[0].peso if ref and not ref.ambigua else None)
        return
    if ref is not None and ref.guias and not ref.ambigua:
        if ref.guias > 1:
            pesos = [c.peso for c in ref.cajas]
            pesos = pesos if len(ref.cajas) == ref.guias else pesos[:1] * ref.guias
            est.fijas += ref.guias
            est.individuales[sku] = est.individuales.get(sku, 0) + ref.guias
            for peso in pesos:
                est.sumar_peso(peso)
            return
        peso = ref.cajas[0].peso
        est.sumar_peso(peso)
        if peso is None:
            est.desconocidas += 1
        elif peso <= PESO_MAXIMO_GUIA:
            est.consolidables.append(peso)
        else:
            est.fijas += 1
        return
    peso_unidad = peso_por_unidad(nombre, sku)
    if es_par_de_pesas(nombre, unidad) and peso_unidad is not None:
        if peso_unidad >= PESO_UNIDAD_GUIA_INDIVIDUAL:
            est.fijas += 2
            est.individuales[sku] = est.individuales.get(sku, 0) + 2
        else:
            est.consolidables.append(peso_unidad * 2)
        est.sumar_peso(peso_unidad * 2)
        return
    est.desconocidas += 1
    est.sumar_peso(None)


def _dentro(valor: float, referencia: float, tolerancia: float, relativa: float = TOLERANCIA_RELATIVA) -> bool:
    return abs(valor - referencia) <= max(tolerancia, referencia * relativa)


def revisar_pedido(pedido: str, lineas: list, referencias: dict[str, Referencia] | None,
                   unidades_picking: dict[str, str], en_revision: bool) -> RevisionGuias:
    """`lineas`: renglones de Liverpool del mismo pedido (LineaLiverpool)."""
    paquetes = lineas[0].paquetes
    observaciones: list[str] = []
    distintos = {(l.paquetes.texto_guias, tuple(l.paquetes.cajas)) for l in lineas}
    if len(distintos) > 1:
        observaciones.append("Los renglones del pedido traen paquetes distintos entre sí en Liverpool; se tomó el primero.")

    est = _Estimacion()
    sin_referencia, ambiguas = [], []
    for linea in lineas:
        ref = (referencias or {}).get(linea.sku)
        if ref is None and referencias is not None:
            sin_referencia.append(linea.sku)
        elif ref is not None and ref.ambigua:
            ambiguas.append(f"{linea.sku} ({ref.nota})")
        for _ in range(max(linea.cantidad, 1)):
            _estimar_unidad(est, linea.sku, linea.producto, unidades_picking.get(linea.sku, ""), ref)

    minimo, maximo = est.minimo, est.maximo
    estimadas = str(minimo) if minimo == maximo else f"{minimo} a {maximo}"
    completa = est.desconocidas == 0

    genericas = [c for c in paquetes.cajas if es_generico(c)]
    if not paquetes.cajas and paquetes.guias in (None, 0):
        observaciones.append("Liverpool no tiene paquetes ni guías registrados para este pedido.")
    elif genericas:
        observaciones.append(
            f"Liverpool tiene {len(genericas)} de {len(paquetes.cajas)} paquetes con medidas genéricas o vacías "
            "(10×10×10 cm, 1 kg o ceros): hay que capturar las medidas reales.")

    guias_lv = paquetes.guias if paquetes.guias not in (None, 0) else (len(paquetes.cajas) or None)
    if guias_lv is not None and paquetes.cajas and guias_lv != len(paquetes.cajas):
        observaciones.append(f"Liverpool indica {guias_lv} guías pero registra {len(paquetes.cajas)} paquetes.")
    if guias_lv is not None and est.fijas + est.piso + len(est.consolidables) > 0:
        if guias_lv < minimo:
            observaciones.append(f"Liverpool solicita {guias_lv} guía(s); según medidas y reglas de empaque se necesitan al menos {minimo}.")
        elif completa and guias_lv > maximo:
            observaciones.append(f"Liverpool solicita {guias_lv} guías; según medidas y reglas de empaque bastan {maximo}.")

    cajas_reales = [c for c in paquetes.cajas if not es_generico(c)]
    referencia_txt = ""
    if len(lineas) == 1 and lineas[0].cantidad == 1:
        ref = (referencias or {}).get(lineas[0].sku)
        if ref and not ref.ambigua and len(ref.cajas) == 1:
            referencia_txt = ref.cajas[0].texto()
            if len(cajas_reales) == 1:
                observaciones.extend(_comparar_caja(cajas_reales[0], ref.cajas[0]))
        elif ref:
            referencia_txt = "; ".join(c.texto() for c in ref.cajas)
    if est.peso_total and cajas_reales and len(cajas_reales) == len(paquetes.cajas) and not referencia_txt:
        peso_lv = sum(c.peso or 0 for c in cajas_reales)
        referencia_txt = f"peso total estimado {_num(est.peso_total)} kg"
        if not _dentro(peso_lv, est.peso_total, TOLERANCIA_PESO_KG, TOLERANCIA_RELATIVA_TOTAL):
            observaciones.append(f"Peso total en Liverpool {_num(peso_lv)} kg; según la tabla de medidas ≈ {_num(est.peso_total)} kg.")

    if est.individuales:
        observaciones.append("Una guía por unidad (pesas de 17 kg o más por pieza o equipo en varias cajas): " + ", ".join(
            f"{sku} → {n} guías" for sku, n in est.individuales.items()) + ".")
    if sin_referencia:
        observaciones.append("Sin medidas de referencia para: " + ", ".join(sorted(set(sin_referencia))) + ".")
    if ambiguas:
        observaciones.append("Referencia ambigua: " + "; ".join(sorted(set(ambiguas))) + ".")
    if referencias is None:
        observaciones.append("No se cargó la tabla de medidas: solo se aplicaron las reglas de empaque.")

    revision = RevisionGuias(
        pedido=pedido,
        productos=", ".join(f"{l.sku}×{l.cantidad}" for l in lineas),
        unidades=sum(l.cantidad for l in lineas),
        en_revision=en_revision,
        guias_liverpool=paquetes.texto_guias or ("" if guias_lv is None else str(guias_lv)),
        guias_estimadas=estimadas + ("" if completa else " (parcial)"),
        paquetes_liverpool="; ".join(c.texto() for c in paquetes.cajas) or "sin paquetes",
        referencia=referencia_txt,
        resultado="CORRECTO",
        observaciones=observaciones,
        minimo=minimo,
        maximo=maximo,
        completa=completa,
    )
    if problemas(revision):
        revision.resultado = "REVISAR"
    return revision


def problemas(revision: RevisionGuias) -> list[str]:
    """Observaciones que requieren acción (excluye notas informativas sobre la referencia)."""
    return [o for o in revision.observaciones if not o.startswith(NOTAS_INFORMATIVAS)]


def _comparar_caja(liverpool: Caja, referencia: Caja) -> list[str]:
    diferencias = []
    if liverpool.peso is not None and referencia.peso is not None and not _dentro(liverpool.peso, referencia.peso, TOLERANCIA_PESO_KG):
        diferencias.append(f"peso {_num(liverpool.peso)} kg en Liverpool vs {_num(referencia.peso)} kg en la tabla")
    if liverpool.medidas and referencia.medidas and not all(
            _dentro(a, b, TOLERANCIA_MEDIDA_CM) for a, b in zip(liverpool.medidas, referencia.medidas)):
        diferencias.append(f"medidas {liverpool.texto().split(' cm')[0]} cm en Liverpool vs "
                           f"{referencia.texto().split(' cm')[0]} cm en la tabla")
    return [f"Medidas de Liverpool no coinciden: {'; '.join(diferencias)}."] if diferencias else []


def revisar_guias(por_pedido: dict[str, list], referencias: dict[str, Referencia] | None,
                  unidades_picking: dict[tuple[str, str], str], pedidos_en_revision: Iterable[str]) -> list[RevisionGuias]:
    en_revision = set(pedidos_en_revision)
    revisiones = []
    for pedido, lineas in por_pedido.items():
        unidades = {sku: u for (p, sku), u in unidades_picking.items() if p == pedido}
        revisiones.append(revisar_pedido(pedido, lineas, referencias, unidades, pedido in en_revision))
    return revisiones


def ordenar(revisiones: list[RevisionGuias]) -> list[RevisionGuias]:
    return sorted(revisiones, key=lambda r: (not r.en_revision, r.resultado != "REVISAR", r.pedido))
