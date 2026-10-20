"""Genera documentos sintéticos (datos inventados) para pruebas.

Uso:
    python tests/generar_ejemplos.py [carpeta_destino]

Crea un Excel de Liverpool, un PDF de picking con anotaciones FreeText, una tabla de medidas de equipos
y un PDF de guías de envío.
Escenario correcto esperado: 0 errores, 4 pedidos verificados, 1 pedido sin picking.
"""

from __future__ import annotations

import io
import re
import sys
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook
from pypdf import PdfReader, PdfWriter
from pypdf.annotations import FreeText
from pypdf.generic import NameObject, TextStringObject
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

ENCABEZADOS_LIVERPOOL = [
    "Fecha de creación", "Shop Id", "Número de proveedor", "Negocio que vende", "Id tienda de venta", "Id del pedido",
    "Estado", "OMS", "Estatus OMS", "Estatus SOMS", "Canal de venta", "Método de envío", "Tipo de artículo",
    "No de SKU", "No de SKU Liverpool", "Nombre del producto", "SKU de la oferta", "Información adicional sku",
    "SKU de seller", "Marca", "Cantidad", "Precio por unidad", "Subtotal", "Valor de la comisión (monto)",
    "Total de la orden en pesos", "Precio original (antes de la promo de liver)", "Precio original de la oferta en pesos",
    "Precio en pesos", "Total de la orden en pesos original", "Total por pieza en pesos", "Método de pago",
    "Promoción aplicada (Liv)", "Moneda", "Pedidos con reembolso", "#Tienda POS", "Negocio que devuelve",
    "Id tienda de devolución", "Reembolso parcial", "Piezas devueltas por indicador", "Reembolso total",
    "Promoción Liverpool", "Nombre completo del cliente que recibe", "Calle 1", "Calle 2", "Colonia", "Municipio",
    "Código postal", "Estado", "Ciudad", "País", "Teléfono principal", "Plazo de envío del pedido",
    "Fecha límite de envío", "Fulfillment (Si/No)", "Código de centro de expedición", "Estado de solicitud de la guía",
    "No de guías requeridas para el envío", "Guías solicitada broker gestor (Si/No)",
    "Guía Solicitada PS (portal del seller) (si ya es solicitada)", "Paquetera para BT: nombre de paquetera",
    "Transportistas SL", "No de guía de entrega", "URL de seguimiento", "Envío asegurado", "Largo del paquete (s)",
    "Alto del paquete (s)", "Ancho del paquete (s)", "Peso del paquete (s)", "Número de contraguías",
    "Fecha de aceptación", "Fecha de envío", "Fecha de entrega", "Fecha de recepción", "Tipo de dispositivo",
    "Versión del aplicativo", "No serie dispositivo (s)", "Subtotal de la linea a pagar al seller",
    "Total de la linea a pagar al seller", "Subtotal de la orden a pagar al seller",
    "Total de la orden a pagar al seller", "solicitud de guia de retorno externa", "Descripcion de SKU",
    "Dimension del sku", "Remision Dummy", "Remision Original",
]


@dataclass
class PedidoFicticio:
    pedido: str
    sku: str
    cantidad: int
    cliente: str
    producto: str = "Producto de prueba"
    estado: str = "Pendiente de envío"
    fecha_limite: str = "22/10/2026 - 06:00 PM"
    reembolso_total: str = "No"
    guia: str = ""
    paquetes: list[tuple[float, float, float, float]] | None = None
    guias: str | None = None

    def columnas_paquetes(self) -> dict[str, str]:
        paquetes = self.paquetes or []
        lista = lambda i: "[" + ", ".join(f"{p[i]:g}" for p in paquetes) + "]"  # noqa: E731
        return {
            "Largo del paquete (s)": lista(0), "Alto del paquete (s)": lista(1), "Ancho del paquete (s)": lista(2),
            "Peso del paquete (s)": lista(3),
            "No de guías requeridas para el envío": self.guias if self.guias is not None else (str(len(paquetes)) if paquetes else ""),
        }


@dataclass
class RenglonFicticio:
    sku: str
    cantidad: float
    unidad: str
    producto: str
    anotaciones: list[str] = field(default_factory=list)
    partir_sku: bool = False


PEDIDOS_BASE = [
    PedidoFicticio("1000000001", "GS-MANC10", 1, "Ana Prueba Ejemplo", "Mancuerna hexagonal 10 kg",
                   fecha_limite="21/10/2026 - 06:00 PM", paquetes=[(30, 20, 30, 20)]),
    PedidoFicticio("1000000002", "GS-BANDALARGA-ROJA", 1, "Bruno Ficticio Demo", "Banda elástica larga roja",
                   paquetes=[(40, 10, 30, 2)]),
    PedidoFicticio("1000000003", "GS-BANDALARGA-ROJA", 2, "Carla Inventada Muestra", "Banda elástica larga roja",
                   paquetes=[(40, 10, 30, 4)]),
    PedidoFicticio("1000000004", "GS-TAPETE", 1, "Diego Sintético Test", "Tapete de yoga negro", guia="GUIA-0001",
                   paquetes=[(60, 15, 15, 2)]),
    PedidoFicticio("1000000005", "GS-BARRA20", 1, "Elena Simulada Caso", "Barra olímpica 20 kg",
                   fecha_limite="20/10/2026 - 01:00 PM", paquetes=[(150, 15, 15, 20)]),
]

MEDIDAS_BASE = [
    ("GS-MANC10", 1, 20, 30, 20, 30),
    ("GS-BANDALARGA-ROJA", 1, 2, 40, 10, 30),
    ("GS-TAPETE", 1, 2, 60, 15, 15),
    ("GS-BARRA20", 1, 20, 150, 15, 15),
]

RENGLONES_BASE = [
    RenglonFicticio("GS-MANC10", 1, "PAR", "Mancuerna hexagonal 10 kg", ["1000000001/ana prueba ejemplo"]),
    RenglonFicticio("GS-BANDALARGA-ROJA", 3, "UNIDAD", "Banda elástica larga roja",
                    ["1000000002/BRUNO FICTICIO DEMO", "1000000003/Carla Inventada Muestra/2 piezas"], partir_sku=True),
    RenglonFicticio("GS-TAPETE", 1, "UNIDAD", "Tapete de yoga negro", ["1000000004/Diego Sintetico Test"]),
]


def crear_excel(destino: Path | io.BytesIO, pedidos: list[PedidoFicticio], *, dimensiones_rotas: bool = False,
                encabezados: list[str] | None = None, hoja: str = "OrderReport") -> None:
    libro = Workbook()
    ws = libro.active
    ws.title = hoja
    encabezados = encabezados or ENCABEZADOS_LIVERPOOL
    ws.append(encabezados)
    for p in pedidos:
        valores = {
            "Fecha de creación": "19/10/2026 - 10:00 AM", "Id del pedido": p.pedido, "Estado": p.estado,
            "Método de envío": "Home delivery", "Nombre del producto": p.producto, "SKU de la oferta": p.sku,
            "Cantidad": str(p.cantidad), "Reembolso parcial": "No", "Reembolso total": p.reembolso_total,
            "Nombre completo del cliente que recibe": p.cliente, "Fecha límite de envío": p.fecha_limite,
            "Estado de solicitud de la guía": "Pendiente", "No de guía de entrega": p.guia,
            **p.columnas_paquetes(),
        }
        fila, estado_usado = [], False
        for encabezado in encabezados:
            if encabezado == "Estado":
                fila.append(p.estado if not estado_usado else "Querétaro")
                estado_usado = True
            else:
                fila.append(valores.get(encabezado, ""))
        ws.append(fila)
        for celda in ws[ws.max_row]:
            if isinstance(celda.value, str) and celda.value.startswith("="):
                celda.data_type = "s"
    buffer = io.BytesIO()
    libro.save(buffer)
    datos = buffer.getvalue()
    if dimensiones_rotas:
        datos = _romper_dimensiones(datos)
    if isinstance(destino, io.BytesIO):
        destino.write(datos)
    else:
        Path(destino).write_bytes(datos)


@dataclass
class GuiaFicticia:
    pedido: str
    contenido: list[str]
    peso: float | None
    rastreo: str
    paqueteria: str = "FedEx"
    medidas: tuple[int, int, int] | None = (30, 30, 20)
    imagen: bool = False


GUIAS_BASE = [
    GuiaFicticia("1000000001", ["1 PAR GS-MANC10"], 20, "8789 0000 0001"),
    GuiaFicticia("1000000002", ["1 PIEZA GS-BANDALARGA-ROJA"], 2, "8789 0000 0002", medidas=(40, 10, 30)),
    GuiaFicticia("1000000003", ["2 PIEZAS GS-BANDALARGA-ROJA"], 4, "8789 0000 0003", medidas=(40, 20, 30)),
    GuiaFicticia("1000000004", ["1 PIEZA GS-TAPETE"], 2, "0000000004", paqueteria="Estafeta", medidas=None),
]


def crear_guias(destino: Path | io.BytesIO, guias: list[GuiaFicticia]) -> None:
    """Guías de envío sintéticas con el formato de texto de FedEx y Estafeta (datos inventados).
    `imagen=True` simula una guía escaneada: solo queda la nota de contenido."""
    buffer = io.BytesIO()
    lienzo = canvas.Canvas(buffer, pagesize=letter)
    alto = letter[1]
    for guia in guias:
        lineas: list[str] = []
        if not guia.imagen and guia.paqueteria == "FedEx":
            lineas = [
                "ORIGIN ID:XXXX 0000000000 SHIP DATE: 20OCT26",
                f"REMITENTE DE PRUEBA ACTWGT: {guia.peso:.2f} KG" if guia.peso is not None else "REMITENTE DE PRUEBA",
                "MARKETPLACE CAD: 000000000/PRUEBA",
                f"CALLE FICTICIA 1 DIMS: {guia.medidas[0]}x{guia.medidas[1]}x{guia.medidas[2]} CM" if guia.medidas else "CALLE FICTICIA 1",
                "TO DESTINATARIO DE PRUEBA",
                f"0000000000 REF: {guia.pedido}Q0PRUEBA01",
                "TRK#",
                guia.rastreo,
            ]
        elif not guia.imagen:
            lineas = [
                "AV. JOSE VASCONCELOS NO. 105, PISO 4, COL. HIPODROMO CONDESA,",
                f"Liverpool {guia.peso:.1f} KG" if guia.peso is not None else "Liverpool",
                f"{guia.pedido}Q0PRUEBA01",
                "Remitente de prueba",
                f"Código de Rastreo: {guia.rastreo}",
                "Vigencia de guía: 19/11/2026",
            ]
        lienzo.setFont("Helvetica", 10)
        for i, linea in enumerate(lineas):
            lienzo.drawString(40, alto - 60 - 16 * i, linea)
        lienzo.setFont("Helvetica-Bold", 14)
        for i, nota in enumerate(guia.contenido):
            lienzo.drawString(300, 120 - 18 * i, nota)
        lienzo.showPage()
    lienzo.save()
    if isinstance(destino, io.BytesIO):
        destino.write(buffer.getvalue())
    else:
        Path(destino).write_bytes(buffer.getvalue())


def crear_medidas(destino: Path | io.BytesIO, filas: list[tuple], *, lado_a_lado: bool = True) -> None:
    """Tabla de medidas como la de Gymsport: EQUIPOS, GUIAS, PESO KG, LARGO, ALTO, ANCHO.
    Con `lado_a_lado`, la mitad de los equipos va en un segundo bloque a la derecha."""
    libro = Workbook()
    ws = libro.active
    ws.title = "INFINITE"
    encabezado = ["EQUIPOS", "GUIAS", "PESO KG", "LARGO CM", "ALTO CM", "ANCHO CM"]
    if lado_a_lado:
        mitad = (len(filas) + 1) // 2
        ws.append(encabezado + encabezado)
        for i in range(mitad):
            derecha = list(filas[mitad + i]) if mitad + i < len(filas) else []
            ws.append(list(filas[i]) + derecha)
    else:
        ws.append(encabezado)
        for fila in filas:
            ws.append(list(fila))
    libro.save(destino)


def _romper_dimensiones(datos: bytes) -> bytes:
    """Reproduce el defecto de Liverpool: la hoja declara una dimensión de 1x1."""
    origen = zipfile.ZipFile(io.BytesIO(datos))
    salida = io.BytesIO()
    with zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED) as destino:
        for entrada in origen.infolist():
            contenido = origen.read(entrada.filename)
            if entrada.filename.startswith("xl/worksheets/sheet"):
                texto = contenido.decode("utf-8")
                if "<dimension" in texto:
                    texto = re.sub(r'<dimension ref="[^"]*"\s*/>', '<dimension ref="A1"/>', texto)
                else:
                    texto = texto.replace("<sheetData", '<dimension ref="A1"/><sheetData', 1)
                contenido = texto.encode("utf-8")
            destino.writestr(entrada, contenido)
    return salida.getvalue()


def crear_picking(destino: Path | io.BytesIO, renglones: list[RenglonFicticio], *, nombre: str = "CONIN/PICK/00001",
                  modo_anotaciones: str = "freetext") -> None:
    """modo_anotaciones: 'freetext' (anotaciones /Annots), 'aplanadas' (texto en otra fuente) o 'ninguna'."""
    base = io.BytesIO()
    lienzo = canvas.Canvas(base, pagesize=letter)
    alto = letter[1]
    posiciones: list[tuple[int, float, list[str]]] = []
    pagina, y = 0, _encabezado(lienzo, nombre, alto, primera=True)

    for renglon in renglones:
        lineas_producto = _envolver(renglon.producto, 18)
        altura = max(70, 14 * (len(lineas_producto) + 2) + 10)
        if y - altura < 90:
            _pie(lienzo, pagina + 1)
            lienzo.showPage()
            pagina += 1
            y = _encabezado(lienzo, nombre, alto, primera=False)
        lienzo.setFont("Helvetica", 9.5)
        if renglon.partir_sku:
            mitad = len(renglon.sku) // 2
            lienzo.drawString(75, y, f"[{renglon.sku[:mitad]}")
            lienzo.drawString(75, y - 14, f"{renglon.sku[mitad:]}]")
            inicio_nombre = y - 28
        else:
            lienzo.drawString(75, y, f"[{renglon.sku}]")
            inicio_nombre = y - 14
        lienzo.drawString(170, y, f"{renglon.cantidad:.2f}")
        lienzo.drawString(170, y - 14, renglon.unidad)
        lienzo.drawString(250, y, "CONIN/Existencias")
        lienzo.drawString(340, y, "CONIN/Existencias/Salida")
        for i, linea in enumerate(lineas_producto):
            lienzo.drawString(75, inicio_nombre - 14 * i, linea)
        top = alto - y
        if modo_anotaciones == "aplanadas":
            lienzo.setFont("Courier-Bold", 12)
            for i, texto in enumerate(renglon.anotaciones):
                lienzo.drawString(215, y - 30 - 14 * i, texto)
        posiciones.append((pagina, top, renglon.anotaciones))
        y -= altura
    _pie(lienzo, pagina + 1)
    lienzo.save()

    lector = PdfReader(io.BytesIO(base.getvalue()))
    escritor = PdfWriter()
    escritor.append(lector)
    if modo_anotaciones == "freetext":
        for numero_pagina, top, anotaciones in posiciones:
            for i, texto in enumerate(anotaciones):
                y_inferior = alto - (top + 36 + 16 * i)
                anotacion = FreeText(text=texto, rect=(215, y_inferior, 560, y_inferior + 14), font="Arial",
                                     bold=True, font_size="12pt", font_color="ff0000", border_color=None,
                                     background_color=None)
                escritor.add_annotation(page_number=numero_pagina, annotation=anotacion)
    salida = io.BytesIO()
    escritor.write(salida)
    if isinstance(destino, io.BytesIO):
        destino.write(salida.getvalue())
    else:
        Path(destino).write_bytes(salida.getvalue())


def modificar_anotacion(origen: bytes, buscar: str, reemplazo: str) -> bytes:
    """Cambia el texto (/Contents) de las anotaciones FreeText que contengan `buscar`."""
    escritor = PdfWriter(clone_from=PdfReader(io.BytesIO(origen)))
    cambios = 0
    for pagina in escritor.pages:
        for ref in pagina.get("/Annots") or []:
            anotacion = ref.get_object()
            contenido = str(anotacion.get("/Contents", ""))
            if anotacion.get("/Subtype") == "/FreeText" and buscar in contenido:
                anotacion[NameObject("/Contents")] = TextStringObject(contenido.replace(buscar, reemplazo))
                cambios += 1
    if not cambios:
        raise ValueError(f"No se encontró ninguna anotación con «{buscar}».")
    salida = io.BytesIO()
    escritor.write(salida)
    return salida.getvalue()


def _encabezado(lienzo, nombre: str, alto: float, primera: bool) -> float:
    lienzo.setFont("Helvetica", 9.5)
    lienzo.drawString(351, alto - 25, "DIRECCIÓN FICTICIA 123")
    lienzo.drawString(351, alto - 40, "00000 CIUDAD DE PRUEBA")
    y = alto - 110
    if primera:
        lienzo.setFont("Helvetica-Bold", 22)
        lienzo.drawString(29, alto - 280, nombre)
        lienzo.setFont("Helvetica", 9.5)
        lienzo.drawString(29, alto - 315, "Orden: Estado: Fecha programada:")
        lienzo.drawString(29, alto - 329, "P00001 Listo 19/10/2026 13:00:00")
        y = alto - 365
    lienzo.setFont("Helvetica-Bold", 9.5)
    for x, texto in ((31, "IMAGEN"), (90, "PRODUCTO"), (164, "CANTIDAD"), (244, "DESDE"), (340, "A"),
                     (437, "CÓDIGO"), (520, "OBSERVACIONES")):
        lienzo.drawString(x, y, texto)
    lienzo.drawString(437, y - 14, "DE BARRAS")
    return y - 45


def _pie(lienzo, numero: int) -> None:
    lienzo.setFont("Helvetica", 9.5)
    lienzo.drawString(214, 70, "DIRECCIÓN FICTICIA 123")
    lienzo.drawString(566, 65, str(numero))


def _envolver(texto: str, ancho: int) -> list[str]:
    lineas, actual = [], ""
    for palabra in texto.split():
        if actual and len(actual) + len(palabra) + 1 > ancho:
            lineas.append(actual)
            actual = palabra
        else:
            actual = f"{actual} {palabra}".strip()
    return lineas + ([actual] if actual else [])


def escenario_correcto(carpeta: Path) -> tuple[Path, Path, Path, Path]:
    carpeta.mkdir(parents=True, exist_ok=True)
    excel = carpeta / "reporte de pedidos 19102026 (sintetico).xlsx"
    pdf = carpeta / "Operaciones de recolección - SINTETICO - CONIN_PICK_00001.pdf"
    medidas = carpeta / "MEDIDAS EQUIPOS (sintetico).xlsx"
    guias = carpeta / "guias de envio (sintetico).pdf"
    crear_excel(excel, PEDIDOS_BASE)
    crear_picking(pdf, RENGLONES_BASE)
    crear_medidas(medidas, MEDIDAS_BASE)
    crear_guias(guias, GUIAS_BASE)
    return excel, pdf, medidas, guias


if __name__ == "__main__":
    destino = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent / "ejemplos_generados"
    excel, pdf, medidas, guias = escenario_correcto(destino)
    print(f"Excel:   {excel}")
    print(f"PDF:     {pdf}")
    print(f"Medidas: {medidas}")
    print(f"Guías:   {guias}")
