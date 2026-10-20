import io
import os
from pathlib import Path

import pytest
from openpyxl import load_workbook

import generar_ejemplos as ej
import verificar_pedidos as vp
from conftest import excel_bytes, pdf_bytes, pedidos_base, renglones_base


def correr(excel=None, pdfs=None):
    excel = excel if excel is not None else excel_bytes()
    pdfs = pdfs if pdfs is not None else [pdf_bytes()]
    return vp.verificar(excel, pdfs)


def problemas(resultado, nivel=None):
    return [h.problema for h in resultado.hallazgos if nivel is None or h.nivel == nivel]


def test_escenario_correcto():
    r = correr()
    assert (r.errores, r.revisar, r.verificados, r.sin_picking) == (0, 0, 4, 1)
    assert r.correctos == 4
    assert r.pedidos_liverpool == 5
    assert r.pedidos_sin_picking[0].pedido == "1000000005"
    assert not r.advertencias
    assert r.todo_cuadra


def test_sku_partido_en_dos_lineas():
    picking = vp.leer_picking(pdf_bytes(), "x.pdf")
    assert [r.sku for r in picking.renglones] == ["GS-MANC10", "GS-BANDALARGA-ROJA", "GS-TAPETE"]
    banda = picking.renglones[1]
    assert banda.cantidad == 3 and banda.unidad == "UNIDAD"
    assert [a.pedido for a in banda.anotaciones] == ["1000000002", "1000000003"]
    assert banda.anotaciones[1].piezas == 2
    assert picking.nombre == "CONIN/PICK/00001"


def test_anotaciones_aplanadas_en_el_texto():
    r = correr(pdfs=[pdf_bytes(modo_anotaciones="aplanadas")])
    assert (r.errores, r.revisar, r.verificados, r.sin_picking) == (0, 0, 4, 1)


def test_regla1_pedido_inexistente_con_sugerencia():
    pdf = ej.modificar_anotacion(pdf_bytes(), "1000000001/", "1000000091/")
    r = correr(pdfs=[pdf])
    errores = problemas(r, vp.ERROR)
    assert any("no existe en Liverpool" in p and "1000000001" in p for p in errores)


def test_regla1_sin_sugerencia_si_no_hay_evidencia():
    pdf = ej.modificar_anotacion(pdf_bytes(), "1000000001/ana prueba ejemplo", "5555599999/persona distinta")
    r = correr(pdfs=[pdf])
    mensaje = next(p for p in problemas(r, vp.ERROR) if "no existe" in p)
    assert "Quisiste decir" not in mensaje and "Posibles" not in mensaje


def test_regla2_sku_incorrecto():
    renglones = renglones_base()
    renglones[2].sku = "GS-TAPETE-AZUL"
    r = correr(pdfs=[pdf_bytes(renglones)])
    assert any("GS-TAPETE-AZUL" in p and "GS-TAPETE" in p for p in problemas(r, vp.ERROR))


def test_regla3_cantidad_del_renglon():
    renglones = renglones_base()
    renglones[1].cantidad = 4
    r = correr(pdfs=[pdf_bytes(renglones)])
    assert any("cantidad del renglón (4)" in p and "(3)" in p for p in problemas(r, vp.ERROR))


def test_regla4_cantidad_en_anotacion():
    pdf = ej.modificar_anotacion(pdf_bytes(), "/2 piezas", "/3 piezas")
    r = correr(pdfs=[pdf])
    assert any("indica 3 piezas" in p and "Liverpool indica 2" in p for p in problemas(r, vp.ERROR))


def test_regla5_nombre_diferente():
    pdf = ej.modificar_anotacion(pdf_bytes(), "Diego Sintetico Test", "Mario Otro Nombre")
    r = correr(pdfs=[pdf])
    assert r.errores == 0
    revisar = problemas(r, vp.REVISAR)
    assert any("Mario Otro Nombre" in p and "Diego Sintético Test" in p for p in revisar)


def test_regla5_normaliza_acentos_mayusculas_y_espacios():
    assert vp.nombres_coinciden("  JOSÉ   pérez ", "jose Perez")
    assert vp.nombres_coinciden("Ma Elena Andrade", "María Elena Andrade López")
    assert not vp.nombres_coinciden("Juan", "Juan Pérez")


def test_regla6_pedido_duplicado_en_mismo_pdf():
    renglones = renglones_base()
    renglones.append(ej.RenglonFicticio("GS-MANC10", 1, "PAR", "Mancuerna", ["1000000001/Ana Prueba Ejemplo"]))
    r = correr(pdfs=[pdf_bytes(renglones)])
    assert any("duplicado dentro del mismo PDF" in p for p in problemas(r, vp.ERROR))


def test_regla6_pedido_duplicado_entre_pdfs():
    segundo = pdf_bytes([ej.RenglonFicticio("GS-TAPETE", 1, "UNIDAD", "Tapete", ["1000000004/Diego Sintetico Test"])],
                        nombre="CONIN/PICK/00002")
    r = vp.verificar(excel_bytes(), [pdf_bytes(), segundo])
    assert any("duplicado entre varios PDF" in p for p in problemas(r, vp.ERROR))


def test_regla6_renglon_sin_pedido():
    renglones = renglones_base()
    renglones.append(ej.RenglonFicticio("GS-BARRA20", 1, "UNIDAD", "Barra olímpica", []))
    r = correr(pdfs=[pdf_bytes(renglones)])
    assert any("Renglón sin número de pedido" in p for p in problemas(r, vp.ERROR))


def test_regla6_pedido_reembolsado():
    pedidos = pedidos_base()
    pedidos[0].reembolso_total = "Si"
    r = correr(excel=excel_bytes(pedidos))
    assert "Pedido reembolsado: no surtir." in problemas(r, vp.ERROR)


def test_regla6_estado_distinto():
    pedidos = pedidos_base()
    pedidos[3].estado = "Pendiente de aceptación"
    r = correr(excel=excel_bytes(pedidos))
    assert r.errores == 0
    assert any("Pendiente de aceptación" in p for p in problemas(r, vp.REVISAR))


def test_regla7_pedidos_sin_picking_ordenados_por_fecha_limite():
    pedidos = pedidos_base()
    pedidos.append(ej.PedidoFicticio("1000000006", "GS-X", 1, "Fulano Prueba", fecha_limite="19/10/2026 - 09:00 AM"))
    r = correr(excel=excel_bytes(pedidos))
    assert [s.pedido for s in r.pedidos_sin_picking] == ["1000000006", "1000000005"]
    assert r.errores == 0


def test_pdf_sin_anotaciones_genera_advertencia():
    r = correr(pdfs=[pdf_bytes(modo_anotaciones="ninguna")])
    assert any("no contiene anotaciones de pedidos reconocibles" in a for a in r.advertencias)
    assert r.verificados == 0
    assert r.errores == 3


def test_excel_con_estructura_incorrecta():
    excel = excel_bytes(encabezados=["Pedido", "Producto", "Unidades"])
    with pytest.raises(vp.DocumentoInvalido, match="no parece ser el reporte de pedidos de Liverpool"):
        correr(excel=excel)


def test_excel_con_dimensiones_internas_incorrectas():
    excel = excel_bytes(dimensiones_rotas=True)
    libro = load_workbook(io.BytesIO(excel), read_only=True)
    assert libro.active.max_row == 1 and libro.active.max_column == 1
    assert len(list(libro.active.iter_rows(values_only=True))) == 1
    libro.close()
    r = correr(excel=excel)
    assert (r.errores, r.verificados, r.sin_picking, r.pedidos_liverpool) == (0, 4, 1, 5)


def test_pdf_danado():
    with pytest.raises(vp.DocumentoInvalido, match="dañado"):
        correr(pdfs=[b"%PDF-1.4\nbasura sin estructura"])


def test_reporte_contiene_hojas_y_evita_formulas(tmp_path):
    pedidos = pedidos_base()
    pedidos[0].cliente = "=HYPERLINK(\"http://malicioso\",\"clic\")"
    r = correr(excel=excel_bytes(pedidos))
    destino = tmp_path / "reporte.xlsx"
    vp.escribir_reporte(r, destino)
    libro = load_workbook(destino)
    assert libro.sheetnames == ["Resumen", "Hallazgos", "Detalle por pedido", "Sin picking", "Guías y medidas"]
    celdas = [c for fila in libro["Detalle por pedido"].iter_rows() for c in fila]
    peligrosas = [c for c in celdas if isinstance(c.value, str) and "HYPERLINK" in c.value]
    assert peligrosas and all(c.data_type == "s" for c in peligrosas)
    assert not any(isinstance(c.value, str) and c.value.startswith("=") for c in celdas)
    assert any(c.value.startswith("'=HYPERLINK") for c in peligrosas)


def test_documentos_reales_opcionales():
    carpeta = Path(os.environ.get("EJEMPLOS_REALES", Path(__file__).resolve().parents[2] / "ejemplos"))
    excels = sorted(carpeta.glob("*.xlsx")) if carpeta.exists() else []
    pdfs = sorted(carpeta.glob("*.pdf")) if carpeta.exists() else []
    if not excels or not pdfs:
        pytest.skip("Sin documentos reales en ejemplos/ (opcional, no se versionan).")
    r = vp.verificar(excels[0], pdfs)
    assert r.pedidos_liverpool > 0
    assert r.detalle
