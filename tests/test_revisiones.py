import copy
import io

import pytest

import generar_ejemplos as ej
import verificar_pedidos as vp
from conftest import Sesion, excel_bytes, pdf_bytes, pedidos_base, renglones_base


def medidas_bytes(filas=None) -> bytes:
    buffer = io.BytesIO()
    ej.crear_medidas(buffer, filas if filas is not None else ej.MEDIDAS_BASE)
    return buffer.getvalue()


def guias_bytes(guias=None) -> bytes:
    buffer = io.BytesIO()
    ej.crear_guias(buffer, guias if guias is not None else ej.GUIAS_BASE)
    return buffer.getvalue()


def guias_base():
    return copy.deepcopy(ej.GUIAS_BASE)


def todos(excel=None, picking=None, medidas=None, guias=None):
    return [
        ("guias.pdf", guias if guias is not None else guias_bytes()),
        ("MEDIDAS.xlsx", medidas if medidas is not None else medidas_bytes()),
        ("CONIN_PICK_00001.pdf", picking if picking is not None else pdf_bytes()),
        ("reporte.xlsx", excel if excel is not None else excel_bytes()),
    ]


def problemas(r, origen=None, nivel=None):
    return [h.problema for h in r.hallazgos if (origen is None or h.origen == origen) and (nivel is None or h.nivel == nivel)]


def test_clasifica_documentos_en_cualquier_orden():
    r = vp.analizar(todos())
    tipos = {d.nombre: d.tipo for d in r.documentos}
    assert tipos == {"reporte.xlsx": vp.TIPO_LIVERPOOL, "MEDIDAS.xlsx": vp.TIPO_MEDIDAS,
                     "CONIN_PICK_00001.pdf": vp.TIPO_PICKING, "guias.pdf": vp.TIPO_GUIAS}
    assert (r.errores, r.revisar, r.verificados, r.correctos, r.sin_picking) == (0, 0, 4, 4, 1)
    assert r.todo_cuadra and r.guias_envio_cargadas and len(r.etiquetas) == 4
    assert all(e.resultado == vp.CORRECTO for e in r.etiquetas)
    assert {g.pedido: g.guias_generadas for g in r.guias}["1000000003"] == "1"


def test_solo_excel_hace_la_primera_revision():
    r = vp.analizar([("reporte.xlsx", excel_bytes()), ("medidas.xlsx", medidas_bytes())])
    assert not r.picking_cargado and not r.guias_envio_cargadas
    assert r.sin_picking is None and r.verificados == 5 and r.correctos == 5
    assert all(g.en_revision for g in r.guias)


def test_solo_reporte_de_liverpool():
    r = vp.analizar([("reporte.xlsx", excel_bytes())])
    assert not r.medidas_cargadas and r.errores == 0


def test_excel_y_guias_sin_picking():
    r = vp.analizar([("reporte.xlsx", excel_bytes()), ("guias.pdf", guias_bytes())])
    assert r.sin_picking is None and r.verificados == 4 and r.errores == 0 and r.revisar == 0


def test_sin_reporte_de_liverpool():
    with pytest.raises(vp.DocumentoInvalido, match="No se encontró el reporte de pedidos de Liverpool"):
        vp.analizar([("picking.pdf", pdf_bytes()), ("guias.pdf", guias_bytes())])


def test_dos_reportes_de_liverpool():
    with pytest.raises(vp.DocumentoInvalido, match="varios reportes de Liverpool"):
        vp.analizar([("a.xlsx", excel_bytes()), ("b.xlsx", excel_bytes())])


def test_documento_no_reconocido_genera_advertencia():
    from reportlab.pdfgen import canvas
    buffer = io.BytesIO()
    lienzo = canvas.Canvas(buffer)
    lienzo.drawString(100, 700, "Factura de compra sin relación")
    lienzo.save()
    r = vp.analizar([("reporte.xlsx", excel_bytes()), ("otro.pdf", buffer.getvalue())])
    assert any("otro.pdf" in a and "no se reconoció" in a for a in r.advertencias)
    assert not r.todo_cuadra


def test_guia_de_pedido_inexistente_con_sugerencia():
    guias = guias_base()
    guias[0].pedido = "1000000091"
    r = vp.analizar(todos(guias=guias_bytes(guias)))
    assert any("no existe en Liverpool" in p and "1000000001" in p for p in problemas(r, vp.ORIGEN_GUIAS, vp.ERROR))


def test_guia_con_sku_ajeno_al_pedido():
    guias = guias_base()
    guias[3].contenido = ["1 PIEZA GS-BARRA20"]
    r = vp.analizar(todos(guias=guias_bytes(guias)))
    assert any("GS-BARRA20" in p and "no pertenece al pedido" in p for p in problemas(r, vp.ORIGEN_GUIAS, vp.ERROR))


def test_guia_duplicada():
    guias = guias_base()
    guias.append(copy.deepcopy(guias[0]))
    r = vp.analizar(todos(guias=guias_bytes(guias)))
    assert any("Guía duplicada" in p for p in problemas(r, vp.ORIGEN_GUIAS, vp.ERROR))
    assert any("Las guías indican 2 unidades de GS-MANC10" in p for p in problemas(r, vp.ORIGEN_GUIAS, vp.ERROR))


def test_peso_de_la_guia_distinto_al_contenido():
    guias = guias_base()
    guias[0].peso = 35
    r = vp.analizar(todos(guias=guias_bytes(guias)))
    assert any("Peso de la guía 35 kg" in p and "≈ 20 kg" in p for p in problemas(r, vp.ORIGEN_GUIAS, vp.REVISAR))


def test_par_pesado_en_una_sola_guia_y_cobertura_por_piezas():
    pedidos = [ej.PedidoFicticio("2000000001", "GS-HEX20", 1, "Ana Prueba", "Set mancuernas de peso fijo 20 kg 2 piezas",
                                 paquetes=[(30, 20, 30, 20), (30, 20, 30, 20)])]
    renglones = [ej.RenglonFicticio("GS-HEX20", 1, "PAR", "Mancuerna 20 kg (PAR)", ["2000000001/Ana Prueba"])]
    excel, picking = excel_bytes(pedidos), pdf_bytes(renglones)

    mal = [ej.GuiaFicticia("2000000001", ["1 PAR GS-HEX20"], 40, "8789 1000 0001")]
    r = vp.analizar(todos(excel, picking, guias=guias_bytes(mal)))
    revisar = problemas(r, vp.ORIGEN_GUIAS, vp.REVISAR)
    assert any("una guía por pieza" in p for p in revisar)
    assert any("Se generaron 1 guías; se necesitan al menos 2" in p for p in revisar)

    bien = [ej.GuiaFicticia("2000000001", ["1 PIEZA GS-HEX20"], 20, f"8789 1000 000{i}") for i in (1, 2)]
    r = vp.analizar(todos(excel, picking, guias=guias_bytes(bien)))
    assert problemas(r, vp.ORIGEN_GUIAS) == []


def test_guia_que_agrupa_mas_de_30_kg():
    medidas = medidas_bytes([("GS-STEP", 1, 12, 90, 15, 35)])
    pedidos = [ej.PedidoFicticio("2000000001", "GS-STEP", 3, "Ana Prueba", "Step", paquetes=[(90, 30, 35, 24), (90, 15, 35, 12)])]
    renglones = [ej.RenglonFicticio("GS-STEP", 3, "UNIDAD", "Step", ["2000000001/Ana Prueba"])]
    guias = [ej.GuiaFicticia("2000000001", ["3 PIEZAS GS-STEP"], 36, "8789 2000 0001")]
    r = vp.analizar(todos(excel_bytes(pedidos), pdf_bytes(renglones), medidas, guias_bytes(guias)))
    assert any("rebasa 30 kg" in p for p in problemas(r, vp.ORIGEN_GUIAS, vp.REVISAR))


def test_guias_que_no_cubren_todas_las_unidades():
    guias = guias_base()
    guias[2].contenido = ["1 PIEZA GS-BANDALARGA-ROJA"]
    guias[2].peso = 2
    r = vp.analizar(todos(guias=guias_bytes(guias)))
    assert any("cubren 1 de 2 unidades de GS-BANDALARGA-ROJA" in p for p in problemas(r, vp.ORIGEN_GUIAS))


def test_pedido_en_picking_sin_guia():
    r = vp.analizar(todos(guias=guias_bytes(guias_base()[:3])))
    assert any("no tiene guía de envío" in h.problema and h.pedido == "1000000004" for h in r.hallazgos)


def test_guia_de_pedido_fuera_del_picking():
    guias = guias_base() + [ej.GuiaFicticia("1000000005", ["1 PIEZA GS-BARRA20"], 20, "8789 0000 0005", medidas=(150, 15, 15))]
    r = vp.analizar(todos(guias=guias_bytes(guias)))
    assert any(h.pedido == "1000000005" and "no aparece en los picking" in h.problema for h in r.hallazgos)


def test_guia_escaneada_se_asigna_por_sku():
    guias = guias_base()
    guias[3] = ej.GuiaFicticia("", ["1 PIEZA GS-TAPETE"], None, "", imagen=True)
    r = vp.analizar(todos(guias=guias_bytes(guias)))
    etiqueta = r.etiquetas[3]
    assert etiqueta.pedido == "1000000004" and etiqueta.identificacion.startswith("Por SKU")
    assert r.errores == 0 and r.revisar == 0


def test_guia_escaneada_ambigua():
    guias = guias_base()[:1] + [ej.GuiaFicticia("", ["1 PIEZA GS-BANDALARGA-ROJA"], None, "", imagen=True)]
    r = vp.analizar(todos(guias=guias_bytes(guias)))
    assert any("No se pudo identificar el pedido" in p and "1000000002" in p and "1000000003" in p
               for p in problemas(r, vp.ORIGEN_GUIAS, vp.REVISAR))


def test_hallazgos_indican_la_revision():
    pedidos = pedidos_base()
    pedidos[0].paquetes = [(10, 10, 10, 1)]
    renglones = renglones_base()
    renglones[2].sku = "GS-TAPETE-AZUL"
    r = vp.analizar(todos(excel_bytes(pedidos), pdf_bytes(renglones)))
    assert {h.origen for h in r.hallazgos} >= {vp.ORIGEN_EXCEL, vp.ORIGEN_PICKING}


def test_reporte_incluye_guias_de_envio(tmp_path):
    from openpyxl import load_workbook
    r = vp.analizar(todos())
    vp.escribir_reporte(r, tmp_path / "r.xlsx")
    libro = load_workbook(tmp_path / "r.xlsx")
    assert "Guías de envío" in libro.sheetnames
    assert libro["Hallazgos"]["B1"].value == "Revisión"


def test_api_con_un_solo_campo_de_documentos(cliente):
    s = Sesion(cliente)
    s.login()
    datos = {"documentos": [(io.BytesIO(contenido), nombre) for nombre, contenido in todos()]}
    r = s.post("/api/verificar", data=datos, content_type="multipart/form-data")
    assert r.status_code == 200
    cuerpo = r.get_json()
    assert [d["tipo"] for d in cuerpo["documentos"]] == [vp.TIPO_LIVERPOOL, vp.TIPO_MEDIDAS, vp.TIPO_PICKING, vp.TIPO_GUIAS]
    assert len(cuerpo["etiquetas"]) == 4 and cuerpo["resumen"]["errores"] == 0

    solo = {"documentos": [(io.BytesIO(excel_bytes()), "reporte.xlsx")]}
    r = s.post("/api/verificar", data=solo, content_type="multipart/form-data")
    assert r.status_code == 200 and r.get_json()["resumen"]["sin_picking"] is None

    nada = {"documentos": [(io.BytesIO(pdf_bytes()), "picking.pdf")]}
    r = s.post("/api/verificar", data=nada, content_type="multipart/form-data")
    assert r.status_code == 422 and "reporte de pedidos de Liverpool" in r.get_json()["error"]


def test_guia_concilia_piezas_del_picking_con_pares_de_liverpool():
    pedidos = [ej.PedidoFicticio("2000000001", "GS-DISCO", 1, "Ana Prueba", "Disco deslizante", paquetes=[(45, 20, 45, 3)])]
    renglones = [ej.RenglonFicticio("GS-DISCO", 2, "UNIDAD", "Disco deslizante", ["2000000001/Ana Prueba/2 piezas"])]
    excel, picking = excel_bytes(pedidos), pdf_bytes(renglones)
    r = vp.analizar([("r.xlsx", excel), ("p.pdf", picking)])
    assert r.errores == 2

    guia = [ej.GuiaFicticia("2000000001", ["2 Piezas (1 Par) GS-DISCO"], 3, "0000000001", paqueteria="Estafeta", medidas=None)]
    r = vp.analizar([("r.xlsx", excel), ("p.pdf", picking), ("g.pdf", guias_bytes(guia))])
    assert r.errores == 0
    conciliados = [h for h in r.hallazgos if "el picking cuenta piezas y Liverpool pares" in h.problema]
    assert len(conciliados) == 2 and all(h.nivel == vp.REVISAR for h in conciliados)
    assert r.detalle[0].resultado == vp.REVISAR
