import io
import re
from pathlib import Path

from openpyxl import load_workbook

from conftest import Sesion
from services.report_service import AlmacenReportes

RAIZ_PROYECTO = Path(__file__).resolve().parents[2]


def test_health(cliente):
    r = cliente.get("/api/health")
    assert r.status_code == 200 and r.get_json() == {"ok": True}
    assert r.headers["Cache-Control"] == "no-store"
    assert r.headers["X-Frame-Options"] == "DENY"
    assert r.headers["Referrer-Policy"] == "no-referrer"


def test_flujo_completo(cliente, documentos):
    s = Sesion(cliente)
    assert s.login().status_code == 200
    excel, pdf = documentos
    r = s.verificar(excel, [("picking.pdf", pdf)])
    assert r.status_code == 200
    datos = r.get_json()
    assert set(datos) >= {"ok", "resumen", "hallazgos", "detalle_pedidos", "pedidos_sin_picking", "advertencias",
                          "reporte_id", "nombre_reporte", "generado", "guias", "medidas_cargadas"}
    assert datos["resumen"] == {"pedidos_liverpool": 5, "verificados": 4, "correctos": 4, "errores": 0,
                                "revisar": 0, "sin_picking": 1, "guias_por_revisar": 0}
    assert datos["medidas_cargadas"] is False and len(datos["guias"]) == 5
    assert re.fullmatch(r"verificacion_pedidos_\d{8}_\d{4}\.xlsx", datos["nombre_reporte"])
    assert len(datos["detalle_pedidos"]) == 4
    assert set(datos["detalle_pedidos"][0]) >= {"resultado", "pedido", "sku_pdf", "cantidad_renglon", "anotacion",
                                                "cliente_liverpool", "sku_liverpool", "cantidad_liverpool", "envio",
                                                "guia", "observaciones"}

    descarga = cliente.get(f"/api/descargar/{datos['reporte_id']}")
    assert descarga.status_code == 200
    assert datos["nombre_reporte"] in descarga.headers["Content-Disposition"]
    assert descarga.headers["Cache-Control"].startswith("no-store") or "no-store" in descarga.headers["Cache-Control"]
    libro = load_workbook(io.BytesIO(descarga.data))
    assert "Hallazgos" in libro.sheetnames

    assert s.post("/api/nueva").status_code == 200
    assert cliente.get(f"/api/descargar/{datos['reporte_id']}").status_code == 404

    assert s.post("/api/auth/logout").status_code == 200
    assert cliente.get("/api/auth/session").status_code == 401


def test_logout_elimina_reporte(app, sesion, documentos):
    excel, pdf = documentos
    assert sesion.verificar(excel, [("p.pdf", pdf)]).status_code == 200
    almacen = app.extensions["verificador"].reportes
    assert len(almacen) == 1
    sesion.post("/api/auth/logout")
    assert len(almacen) == 0


def test_nueva_verificacion_reemplaza_reporte_anterior(app, sesion, documentos):
    excel, pdf = documentos
    primero = sesion.verificar(excel, [("p.pdf", pdf)]).get_json()["reporte_id"]
    segundo = sesion.verificar(excel, [("p.pdf", pdf)]).get_json()["reporte_id"]
    assert primero != segundo
    assert len(app.extensions["verificador"].reportes) == 1
    assert sesion.cliente.get(f"/api/descargar/{primero}").status_code == 404
    assert sesion.cliente.get(f"/api/descargar/{segundo}").status_code == 200


def test_reporte_ajeno_no_se_puede_descargar(app, documentos):
    a, b = Sesion(app.test_client()), Sesion(app.test_client())
    a.login()
    b.login()
    excel, pdf = documentos
    reporte = a.verificar(excel, [("p.pdf", pdf)]).get_json()["reporte_id"]
    assert b.cliente.get(f"/api/descargar/{reporte}").status_code == 404
    assert a.cliente.get(f"/api/descargar/{reporte}").status_code == 200


def test_reporte_expirado(app, sesion, documentos):
    ahora = [0.0]
    app.extensions["verificador"].reportes = AlmacenReportes(3600, 50, 100 * 1024 * 1024, reloj=lambda: ahora[0])
    excel, pdf = documentos
    reporte = sesion.verificar(excel, [("p.pdf", pdf)]).get_json()["reporte_id"]
    ahora[0] = 3599
    assert sesion.cliente.get(f"/api/descargar/{reporte}").status_code == 200
    ahora[0] = 3601
    r = sesion.cliente.get(f"/api/descargar/{reporte}")
    assert r.status_code == 404 and "Repite la verificación" in r.get_json()["error"]


def test_identificador_invalido(sesion):
    assert sesion.cliente.get("/api/descargar/../../etc").status_code == 404
    assert sesion.cliente.get("/api/descargar/corto").status_code == 404


def test_almacen_respeta_limites_de_memoria():
    almacen = AlmacenReportes(3600, max_reportes=2, max_bytes=10)
    almacen.guardar("a", b"12345", "a.xlsx")
    almacen.guardar("b", b"12345", "b.xlsx")
    almacen.guardar("c", b"12345", "c.xlsx")
    assert len(almacen) == 2 and almacen.bytes_usados <= 10


def test_cors_solo_origenes_permitidos(directorio_temporal):
    from conftest import configuracion
    from factory import create_app

    app = create_app(configuracion(cors_origins=("https://verificador.example.com",)),
                     directorio_temporal=str(directorio_temporal), limpieza_periodica=False)
    cliente = app.test_client()
    permitido = cliente.get("/api/health", headers={"Origin": "https://verificador.example.com"})
    assert permitido.headers["Access-Control-Allow-Origin"] == "https://verificador.example.com"
    assert permitido.headers["Access-Control-Allow-Credentials"] == "true"
    ajeno = cliente.get("/api/health", headers={"Origin": "https://malicioso.example.com"})
    assert "Access-Control-Allow-Origin" not in ajeno.headers
    previa = cliente.options("/api/verificar", headers={"Origin": "https://verificador.example.com",
                                                        "Access-Control-Request-Method": "POST"})
    assert previa.status_code == 200


def test_no_existe_base_de_datos_ni_dependencias_de_bd():
    requisitos = (RAIZ_PROYECTO / "backend" / "requirements.txt").read_text(encoding="utf-8").lower()
    for prohibido in ("sqlalchemy", "psycopg", "pymongo", "redis", "sqlite", "prisma", "supabase", "firebase"):
        assert prohibido not in requisitos
    extensiones = {".db", ".sqlite", ".sqlite3"}
    encontrados = [p for p in (RAIZ_PROYECTO / "backend").rglob("*")
                   if p.suffix in extensiones and ".venv" not in p.parts]
    assert encontrados == []


def test_no_se_guardan_documentos_en_frontend(sesion, documentos):
    publico = RAIZ_PROYECTO / "frontend" / "public"
    antes = sorted(p.name for p in publico.rglob("*")) if publico.exists() else []
    excel, pdf = documentos
    assert sesion.verificar(excel, [("p.pdf", pdf)]).status_code == 200
    despues = sorted(p.name for p in publico.rglob("*")) if publico.exists() else []
    assert antes == despues
